"""Printer template arithmetic: no eval, VBA, live writes or inferred shop IDs."""
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from typing import Literal
from pydantic import Field, model_validator
from services.channelmanager import Strict, Semantics, number

LABELS = ("123inkt", "Printerland", "Printabout", "Sneltoner", "Tonershop")
SHEETS = ("Prijzen pre wijziging", "Prijs adviezen PriceShape", "Formules",
          "Formules inplakken", "Prijzen na wijziging")


class ConditionalRule(Strict):
    threshold: str = "50"
    lower_offset: str
    upper_offset: str
    multiplier: str = "1"
    round_inner_integer: bool = False
    lower_tail: str = "0"
    upper_tail: str = "0"

    @model_validator(mode="after")
    def valid(self):
        for field in ("threshold", "lower_offset", "upper_offset", "multiplier", "lower_tail", "upper_tail"):
            if number(getattr(self, field)) < 0:
                raise ValueError("Printer conditional parameters cannot be negative")
        return self


def default_rules():
    return {
        "Printerland": ConditionalRule(lower_offset="6", upper_offset="5"),
        "Printabout": ConditionalRule(lower_offset="1.50", upper_offset=".50"),
        "Sneltoner": ConditionalRule(lower_offset="3", upper_offset="2"),
        "Tonershop": ConditionalRule(lower_offset="2", upper_offset="1", multiplier="1.1",
                                     round_inner_integer=True, lower_tail="2.99", upper_tail="1.99")}


class PrinterConfig(Semantics):
    active: bool = False
    mapping_confirmed: bool = False
    shops: dict[str, int]
    id_transform: Literal["exact", "nl_six_digits"] = "exact"
    transformation_confirmed: bool = False
    rules: dict[str, ConditionalRule] = Field(default_factory=default_rules)
    export_columns: dict[str, Literal["sku", "shop_id", "label", "target_price"]] = Field(
        default_factory=lambda: {"Productcode": "sku", "Shop": "shop_id", "Prijs": "target_price"})

    @model_validator(mode="after")
    def valid(self):
        if "123inkt" not in self.shops or set(self.shops) - set(LABELS):
            raise ValueError("Explicit 123inkt mapping required; unknown printer labels forbidden")
        if set(self.rules) - set(LABELS[1:]) or (set(self.shops) - {"123inkt"}) - set(self.rules):
            raise ValueError("Every mapped target requires an explicit supported conditional rule")
        if any(v <= 0 for v in self.shops.values()) or len(set(self.shops.values())) != len(self.shops):
            raise ValueError("Shop IDs must be positive and unique")
        if self.active and not self.mapping_confirmed:
            raise ValueError("Confirm label-to-shop mapping before activation")
        if self.id_transform != "exact" and not self.transformation_confirmed:
            raise ValueError("Confirm explicit product ID transformation")
        if not {"sku", "shop_id", "target_price"} <= set(self.export_columns.values()):
            raise ValueError("Paste table requires sku, shop_id and target_price")
        if len(self.export_columns) > 10 or any(not h or len(h) > 100 for h in self.export_columns):
            raise ValueError("Invalid paste table headers")
        return self


class PriceRow(Strict):
    sku: str = Field(min_length=1, max_length=500)
    shop_id: int = Field(gt=0)
    price: str

    @model_validator(mode="after")
    def valid(self):
        if number(self.price) < 0:
            raise ValueError("Negative current prices are invalid")
        return self


def money(value):
    return str(number(value).quantize(Decimal(".01"), rounding=ROUND_HALF_UP))


def transformed(sku, config):
    if config.id_transform == "exact":
        return sku
    if not sku.isascii() or not sku.isdigit() or len(sku) > 6:
        raise ValueError("nl_six_digits only accepts 1–6 ASCII numeric source IDs")
    return "nl" + sku.zfill(6)


def price_index(rows, shops):
    indexed, duplicates = {}, 0
    for raw in rows:
        row = PriceRow.model_validate(raw)
        if row.shop_id not in shops:
            raise ValueError("Price row shop outside explicit mapping")
        key = (row.sku, row.shop_id)
        value = number(row.price)
        if key in indexed:
            if indexed[key] != value:
                raise ValueError("Conflicting prices for product/shop")
            duplicates += 1
        indexed[key] = value
    if not indexed:
        raise ValueError("Empty price snapshot")
    return indexed, duplicates


def target_price(label, p, rules=None):
    if label == "123inkt":
        result = p
    else:
        rule = (rules if rules is not None else default_rules()).get(label)
        if rule is None:
            raise ValueError("Unknown printer label")
        low = p < number(rule.threshold)
        result = (p + number(rule.lower_offset if low else rule.upper_offset)) * number(rule.multiplier)
        if rule.round_inner_integer:
            result = result.quantize(Decimal("1"), rounding=ROUND_HALF_UP)
        result += number(rule.lower_tail if low else rule.upper_tail)
    return money(result)


def calculate(config, advice, pre):
    indexed, duplicates = price_index(pre, set(config.shops.values()))
    products, advice_duplicates = {}, 0
    source_ids = {}
    for item in advice:
        source = str(item["productcode"])
        sku = transformed(source, config)
        p = number(item["advice"])
        if p < 0:
            raise ValueError("Negative advice")
        if sku in products:
            if products[sku] != p or source_ids[sku] != source:
                raise ValueError("Conflicting advice or transformed product ID collision")
            advice_duplicates += 1
        products[sku], source_ids[sku] = p, source
    if not products:
        raise ValueError("Empty advice snapshot")
    rows, warnings, blocks = [], [], []
    for sku, p in products.items():
        old = indexed.get((sku, config.shops["123inkt"]))
        reason = None
        if old is None:
            reason = "missing_basis_pre_price"
        elif old == 0:
            reason = "zero_basis_comparison_price"
        elif p < 20:
            reason = "advice_below_20_do_not_adjust"
        elif abs((p - old) / old) > Decimal(".20"):
            warnings.append({"sku": sku, "reason": "change_over_20_percent"})
        for label, shop in config.shops.items():
            current = indexed.get((sku, shop))
            state = reason
            if state is None and current is None:
                state = "missing_target_pre_price"
            if state is None and current == 0:
                state = "zero_target_pre_price_suppressed"
            if state and state != "zero_target_pre_price_suppressed":
                blocks.append({"sku": sku, "shop_id": shop, "reason": state})
            rows.append({"sku": sku, "source_id": source_ids[sku], "shop_id": shop,
                         "label": label, "advice": money(p),
                         "current_price": money(current) if current is not None else None,
                         "target_price": target_price(label, p, config.rules) if state is None else None,
                         "state": state or "ready"})
    return {"rows": rows, "warnings": warnings, "blocks": blocks,
            "pre_duplicates": duplicates, "advice_duplicates": advice_duplicates,
            "pre_only_products": sorted({s for s, _ in indexed} - set(products))}


def reconcile(export_rows, actual_rows, shops):
    """Full-shop readback: only exported non-suppressed targets are assertions.

    Other products remain visible as out-of-scope rows, not price failures.
    All rows must still belong to authorized mapped shops and have valid prices.
    """
    actual, duplicates = price_index(actual_rows, shops)
    expected = {(r["sku"], r["shop_id"]): number(r["target_price"])
                for r in export_rows if r["target_price"] is not None}
    checks = [{"sku": sku, "shop_id": shop, "expected": money(value),
               "actual": money(actual[(sku, shop)]) if (sku, shop) in actual else None,
               "difference": money(value - actual[(sku, shop)]) if (sku, shop) in actual else None,
               "state": ("missing" if (sku, shop) not in actual else
                         "matched" if actual[(sku, shop)] == value else "mismatch")}
              for (sku, shop), value in expected.items()]
    extras = [{"sku": sku, "shop_id": shop} for sku, shop in actual.keys() - expected.keys()]
    return {"checks": checks, "extra_rows": extras, "duplicates": duplicates,
            "scope": "exported_targets_in_full_snapshot",
            "confirmed": bool(checks) and all(c["state"] == "matched" for c in checks)}


def safe_text(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def workbook(config, calculation, provenance, pre, advice, readback=None):
    from openpyxl import Workbook
    book = Workbook()
    book.remove(book.active)
    for title in SHEETS:
        sheet = book.create_sheet(title)
        sheet.append(["Printer workflow — DRS-compatibiliteit onbevestigd"])
        for key, value in provenance.items():
            sheet.append([safe_text(str(key)), safe_text(str(value))])
        sheet.append([])
        sheet.freeze_panes = f"A{sheet.max_row + 2}"
    def table(sheet, headers, values):
        sheet.append([safe_text(v) for v in headers])
        for row in values:
            sheet.append([safe_text(v) for v in row])
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 25
    table(book[SHEETS[0]], ["Productcode", "Shop ID", "Pre prijs"],
          [[r["sku"], r["shop_id"], r["price"]] for r in pre])
    basis_rows = {r["source_id"]: r for r in calculation["rows"] if r["label"] == "123inkt"}
    warned = {r["sku"] for r in calculation["warnings"]}
    advice_values = []
    for r in advice:
        basis = basis_rows[r["productcode"]]
        old = basis["current_price"]
        change = str((number(r["advice"]) - number(old)) / number(old) * 100) if old and number(old) else None
        check = ("change_over_20_percent" if basis["sku"] in warned else basis["state"])
        advice_values.append([r["productcode"], old, r["advice"], change, check])
    table(book[SHEETS[1]], ["External product ID", "Old basis price", "Approved advice", "Change %", "Check"],
          advice_values)
    fields = ["sku", "source_id", "shop_id", "label", "current_price", "advice", "target_price", "state"]
    table(book[SHEETS[2]], fields, [[r.get(f) for f in fields] for r in calculation["rows"]])
    table(book[SHEETS[3]], list(config.export_columns),
          [[r[f] for f in config.export_columns.values()] for r in calculation["rows"]
           if r["target_price"] is not None])
    controls = ([[r["sku"], r["shop_id"], r["expected"], r["actual"], r["difference"], r["state"]]
                 for r in readback["checks"]] if readback is not None else
                [[r["sku"], r["shop_id"], r["target_price"], None, None, "unknown_not_read_back"]
                 for r in calculation["rows"] if r["target_price"] is not None])
    table(book[SHEETS[4]], ["Productcode", "Shop ID", "Expected", "Actual", "Difference", "Status"], controls)
    if readback:
        for extra in readback["extra_rows"]:
            book[SHEETS[4]].append([safe_text(extra["sku"]), extra["shop_id"], None, None, None, "unexpected_row"])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
