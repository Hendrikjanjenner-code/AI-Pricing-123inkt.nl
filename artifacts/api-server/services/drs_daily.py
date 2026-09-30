"""Two-header DRS daily sales: exact signed values, staged and atomic.

File bytes are never persisted. Parsed product/day arrays are stored in bounded
JSON chunks. Excel is read-only streamed; lookups and writes are batched.
"""
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from uuid import uuid4
import re
import csv
import io
import copy
import json
import openpyxl
from fastapi import HTTPException
from sqlalchemy import tuple_
import models
from services.product_identity import IdentityIndex, sales_lock, alias_discriminator, ean_values

TEMPLATE = "drs-sales-wide-v1"
DEFAULT_MAPPING = {"sheet": "sales", "code_column": 1, "ean_column": 3, "first_day_column": 5,
                   "name_column": 2, "oem_column": 4}
WARNINGS = [
    "Retourverwerking niet geverifieerd: exact getekende aantallen en bedragen, geen extra retourcorrectie.",
    "Bedragen in EUR; analyseomzet exclusief btw, inclusief btw apart bewaard.",
    "Historische kosten/marge onbekend; huidige inkoopprijs en productmarge worden niet gebruikt.",
    "Nulwaarden zijn expliciete dagcorrecties. Andere winkels, bronnen en datums blijven intact.",
]


def bulk_daily_insert(db, model, rows):
    """Bounded PostgreSQL COPY on the existing transaction, not per-row SQL.

    Fixed model/column names only; CSV escaping protects all values. Database
    constraints and transaction rollback remain authoritative. SQLite tests can
    still use SQLAlchemy's bulk implementation.
    """
    if not rows:
        return
    if db.bind.dialect.name != "postgresql":
        db.bulk_insert_mappings(model, rows)
        return
    columns = list(rows[0])
    if model not in (models.SalesDailyClaim, models.FactSalesDaily):
        raise ValueError("Unsupported daily bulk model")
    if any(c not in model.__table__.columns for c in columns):
        raise ValueError("Unknown daily bulk column")
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerows([[r.get(c) for c in columns] for r in rows])
    buffer.seek(0)
    quoted = ", ".join('"' + c + '"' for c in columns)
    raw = db.connection().connection.driver_connection
    with raw.cursor() as cursor:
        cursor.copy_expert(f'COPY "{model.__tablename__}" ({quoted}) FROM STDIN WITH (FORMAT CSV)', buffer)


def identifier(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        if int(value) != value:
            raise ValueError("Artikelidentificatie moet gehele tekst zijn.")
        text = str(int(value))
        if re.fullmatch(r"0+", cell.number_format or ""):
            text = text.zfill(len(cell.number_format))
        return text
    return str(value).strip()


def amount(value):
    if value is None or value == "":
        return Decimal(0)
    try:
        n = Decimal(str(value))
    except InvalidOperation:
        raise ValueError(f"Ongeldig bedrag/aantal: {value!r}")
    if not n.is_finite() or abs(n) >= Decimal("100000000"):
        raise ValueError("Bedrag/aantal niet eindig of buiten databasebereik.")
    rounded = n.quantize(Decimal(".0001"))
    if abs(n - rounded) > Decimal(".00000001"):
        raise ValueError("Meer dan vier decimalen; geen stille afronding toegestaan.")
    return rounded


def parse_workbook(file, year, start, end, mapping=None):
    """Return compact rows and validation errors; errors cap is display-only."""
    m = {**DEFAULT_MAPPING, **(mapping or {})}
    if set(m) != set(DEFAULT_MAPPING):
        raise ValueError("Onbekende kolomkoppeling.")
    if not 1900 <= year <= 2200 or end < start or (end - start).days > 366:
        raise ValueError("Ongeldig jaar of periode (maximaal 366 dagen).")
    for key in ("code_column", "ean_column", "first_day_column", "name_column", "oem_column"):
        if not isinstance(m[key], int) or not 1 <= m[key] <= 16384:
            raise ValueError("Kolommen moeten 1-based Excel-kolomnummers zijn.")
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    try:
        names = {n.lower(): n for n in wb.sheetnames}
        if str(m["sheet"]).lower() not in names:
            raise ValueError("Sales-tabblad ontbreekt.")
        ws = wb[names[str(m["sheet"]).lower()]]
        rows = ws.iter_rows()
        headings, headers = next(rows), next(rows)
        first = m["first_day_column"] - 1
        dates, columns = [], []
        current_year, previous = year, None
        total_col = None
        for col in range(first, len(headings), 3):
            raw = headings[col].value
            if str(raw).strip().lower() in ("totaal", "total"):
                total_col = col
                break
            if raw is None:
                raise ValueError(f"Ontbrekende dagkop in kolom {col + 1}.")
            if isinstance(raw, (datetime, date)):
                day, month = raw.day, raw.month
            else:
                match = re.fullmatch(r"\s*(\d{1,2})/(\d{1,2})\s*", str(raw))
                if not match:
                    raise ValueError(f"Ongeldige dagkop: {raw!r}.")
                day, month = map(int, match.groups())
            if previous and (month, day) < previous:
                if previous[0] != 12 or month != 1:
                    raise ValueError("Dagen staan niet op volgorde.")
                current_year += 1
            d = date(current_year, month, day)
            if dates and d != date.fromisoformat(dates[-1]) + timedelta(days=1):
                raise ValueError("Dagen moeten aaneengesloten en uniek zijn.")
            previous = (month, day)
            actual = [str(c.value or "").strip().lower().replace(".", "")
                      for c in headers[col:col + 3]]
            if actual != ["quantity", "revenue incl vat", "revenue excl vat"]:
                raise ValueError(f"Onverwachte dagkolommen bij {raw!r}: {actual}.")
            dates.append(d.isoformat())
            columns.append(col)
        if not dates or total_col is None:
            raise ValueError("Geen dagkolommen of Totaal-controle gevonden.")
        if date.fromisoformat(dates[0]) != start or date.fromisoformat(dates[-1]) != end:
            raise ValueError("Opgegeven periode wijkt af van dagkoppen en expliciet jaar.")
        parsed, errors, count = [], [], 0
        day_sums = [[Decimal(0), Decimal(0), Decimal(0)] for _ in dates]
        footer = None
        totals = [Decimal(0), Decimal(0), Decimal(0)]
        seen = set()
        for rownum, cells in enumerate(rows, 3):
            code, ean, name, oem = "", "", "", ""
            try:
                code = identifier(cells[m["code_column"] - 1])
                if not code:
                    if str(cells[1].value or "").strip().lower() in ("totaal", "total"):
                        if footer is not None:
                            raise ValueError("Meerdere totaalsluitregels.")
                        footer = cells
                        continue
                    if any(c.value is not None for c in cells):
                        raise ValueError("Artikelcode ontbreekt.")
                    continue
                count += 1
                ean = identifier(cells[m["ean_column"] - 1])
                name = identifier(cells[m["name_column"] - 1])
                oem = identifier(cells[m["oem_column"] - 1])
                if len(code) > 100 or len(ean) > 4096:
                    raise ValueError("Identificatie te lang; niet afgekapt.")
                if (code, ean) in seen:
                    raise ValueError("Dubbele artikelcode/EAN in bestand.")
                seen.add((code, ean))
                values, sums = [], [Decimal(0), Decimal(0), Decimal(0)]
                for col in columns:
                    triple = [amount(c.value) for c in cells[col:col + 3]]
                    if len(triple) != 3 or triple[0] != triple[0].to_integral_value():
                        raise ValueError("Aantal moet een geheel getekend aantal zijn.")
                    values.append([int(triple[0]), str(triple[1]), str(triple[2])])
                    sums = [a + b for a, b in zip(sums, triple)]
                controls = [amount(c.value) for c in cells[total_col:total_col + 3]]
                if len(controls) != 3 or any(abs(a - b) > Decimal(".0001") for a, b in zip(sums, controls)):
                    raise ValueError("Dagtotalen wijken af van Totaal-kolommen.")
                totals = [a + b for a, b in zip(totals, sums)]
                for i, triple in enumerate(values):
                    day_sums[i] = [a + Decimal(str(b)) for a, b in zip(day_sums[i], triple)]
                parsed.append({"row": rownum, "code": code, "ean": ean,
                               "product_name": name, "oem": oem, "values": values})
            except (ValueError, IndexError) as exc:
                errors.append({"row": rownum, "code": code, "ean": ean, "product_name": name,
                               "oem": oem, "message": str(exc), "kind": "parse", "candidates": []})
        if not count:
            errors.append({"row": None, "code": "", "message": "Geen artikelregels gevonden."})
        if footer is not None and not errors:
            for i, col in enumerate(columns):
                controls = [amount(c.value) for c in footer[col:col + 3]]
                if controls != day_sums[i]:
                    errors.append({"row": None, "code": "", "message": f"Totaalsluitregel wijkt af op {dates[i]}."})
        return parsed, dates, errors, {
            "quantity": int(totals[0]), "revenue_incl_vat": str(totals[1]),
            "revenue_excl_vat": str(totals[2]),
        }
    finally:
        wb.close()


def create_preview(db, user, shop_id, file, year, start, end, mapping):
    rows, dates, errors, totals = parse_workbook(file, year, start, end, mapping)
    return stage_rows(db, user, shop_id, rows, dates, errors, totals)


def stage_rows(db, user, shop_id, rows, dates, errors, totals, previous=None):
    start, end = date.fromisoformat(dates[0]), date.fromisoformat(dates[-1])
    parse_errors = copy.deepcopy(errors)
    index = IdentityIndex(db, shop_id)
    proposals, sample, resolved = [], [], set()
    for row in rows:
        row.pop("sku_id", None)
        try:
            sku, proposed = index.resolve(row["code"], row["ean"])
            if sku in resolved:
                raise ValueError("Meerdere bestandsregels wijzen naar hetzelfde product.")
            resolved.add(sku)
            row["sku_id"] = sku
            if proposed:
                proposals.append({k: row[k] for k in ("code", "ean", "sku_id")})
            if len(sample) < 10:
                for d, v in zip(dates, row["values"]):
                    if any(Decimal(str(x)) for x in v):
                        sample.append({"sku_id": sku, "business_sku": row["code"], "date": d,
                                       "quantity": v[0], "revenue_incl_vat": v[1], "revenue_excl_vat": v[2]})
                        break
        except ValueError as exc:
            errors.append({k: row.get(k, "") for k in ("row", "code", "ean", "product_name", "oem")} |
                          {"message": str(exc), "kind": "identity", "candidates": index.evidence(row)})
    # Legacy/unowned overlap is never deleted or silently adopted.
    legacy = db.query(models.FactSalesDaily.sku_id, models.FactSalesDaily.date).filter(
        models.FactSalesDaily.shop_id == shop_id,
        models.FactSalesDaily.date.between(start, end),
        (models.FactSalesDaily.source.is_(None)) | (models.FactSalesDaily.source != "DRS_DAILY"),
    )
    for sku, d in legacy:
        if sku in resolved:
            errors.append({"row": None, "code": sku, "message": f"Bestaande andere/legacy verkoopbron op {d}."})
    stage_id = str(uuid4())
    expires = datetime.utcnow() + timedelta(hours=24)
    preview = {"stage_id": stage_id, "status": "blocked" if errors else "ready",
               "shop_id": shop_id, "period_start": start.isoformat(), "period_end": end.isoformat(),
               "product_count": len(rows), "fact_count": len(rows) * len(dates), "totals": totals,
               "errors": errors, "error_count": len(errors), "warnings": WARNINGS,
               "metadata_version": 2, "parse_errors": parse_errors,
               "mapping_proposals": proposals[:200], "mapping_proposal_count": len(proposals),
               "sample": sample, "expires_at": expires.isoformat(), "dates": dates}
    if previous:
        # A changed error message is not a resolved row.
        old = {(e.get("row"), e.get("code")) for e in previous.preview["errors"]}
        new = {(e.get("row"), e.get("code")) for e in errors}
        preview.update(previous_stage_id=previous.id, previous_error_count=len(old),
                       resolved_error_count=len(old - new))
    stage = models.DrsDailyStage(id=stage_id, shop_id=shop_id, created_by=user.username,
                                expires_at=expires, status=preview["status"], preview=preview)
    db.add(stage)
    db.flush()
    for offset in range(0, len(rows), 100):
        db.add(models.DrsDailyChunk(stage_id=stage_id, rows=rows[offset:offset + 100]))
    db.commit()
    return preview


def confirm_stage(db, stage, confirm_mapping):
    sales_lock(db, stage.shop_id)
    # Serialize repeat claims and reload after waiting on a concurrent commit.
    stage = db.query(models.DrsDailyStage).filter_by(id=stage.id).with_for_update().populate_existing().one()
    if stage.status == "confirmed":
        return stage.result
    if stage.expires_at < datetime.utcnow():
        raise HTTPException(410, "Preview verlopen; maak een nieuwe preview.")
    if stage.status != "ready":
        raise HTTPException(409, "Geblokkeerde preview kan niet worden bevestigd.")
    if stage.preview["mapping_proposal_count"] and not confirm_mapping:
        raise HTTPException(409, "Bevestig expliciet de voorgestelde artikelcode-mapping.")
    index = IdentityIndex(db, stage.shop_id)
    dates = [date.fromisoformat(d) for d in stage.preview["dates"]]
    result = {"stage_id": stage.id, "status": "confirmed", "inserted": 0, "updated": 0, "unchanged": 0}
    chunks = db.query(models.DrsDailyChunk).filter_by(stage_id=stage.id).order_by(models.DrsDailyChunk.id)
    for chunk in chunks.yield_per(1):
        sku_ids = [r["sku_id"] for r in chunk.rows]
        existing = {}
        for f in db.query(models.FactSalesDaily).filter(
            models.FactSalesDaily.shop_id == stage.shop_id,
            models.FactSalesDaily.sku_id.in_(sku_ids), models.FactSalesDaily.date.in_(dates)):
            if f.source != "DRS_DAILY" or (f.sku_id, f.date) in existing:
                raise HTTPException(409, "Overlap met bestaande andere/legacy verkoopbron.")
            existing[(f.sku_id, f.date)] = f
        claims = {(c.sku_id, c.date): c for c in db.query(models.SalesDailyClaim).filter(
            models.SalesDailyClaim.shop_id == stage.shop_id,
            models.SalesDailyClaim.sku_id.in_(sku_ids), models.SalesDailyClaim.date.in_(dates))}
        inserts, updates, new_claims = [], [], []
        for row in chunk.rows:
            try:
                sku, proposed = index.resolve(row["code"], row["ean"])
            except ValueError as exc:
                raise HTTPException(409, f"Mapping gewijzigd: {exc}")
            if sku != row["sku_id"]:
                raise HTTPException(409, "Mapping gewijzigd sinds preview.")
            if proposed:
                db.add(models.ProductSourceAlias(shop_id=stage.shop_id, source="DRS",
                       code=row["code"], discriminator=alias_discriminator(row["ean"]), sku_id=sku))
                index.aliases[(row["code"], alias_discriminator(row["ean"]))].add(sku)
            if not index.products[sku].business_sku:
                index.products[sku].business_sku = row["code"]
            for d, values in zip(dates, row["values"]):
                key = (sku, d)
                claim = claims.get(key)
                if claim and claim.source != "DRS_DAILY":
                    raise HTTPException(409, "Dag is al geclaimd door een andere verkoopbron.")
                if not claim:
                    new_claims.append({"shop_id": stage.shop_id, "sku_id": sku, "date": d, "source": "DRS_DAILY"})
                units, incl, excl = values[0], Decimal(values[1]), Decimal(values[2])
                fact = existing.get(key)
                data = {"units_sold": units, "revenue": excl, "revenue_incl_vat": incl,
                        "cost": None, "margin": None, "margin_pct": None}
                if fact:
                    if fact.units_sold == units and fact.revenue == excl and fact.revenue_incl_vat == incl:
                        result["unchanged"] += 1
                    else:
                        updates.append({"id": fact.id, **data})
                        result["updated"] += 1
                elif units or incl or excl:
                    inserts.append({"date": d, "sku_id": sku, "shop_id": stage.shop_id,
                                    "source": "DRS_DAILY", **data})
                    result["inserted"] += 1
                else:
                    result["unchanged"] += 1
        # All writes share one transaction. The unique daily claim is the final
        # database-level concurrency guard, independently of HTTP retries.
        if new_claims:
            bulk_daily_insert(db, models.SalesDailyClaim, new_claims)
        if inserts:
            bulk_daily_insert(db, models.FactSalesDaily, inserts)
        if updates:
            db.bulk_update_mappings(models.FactSalesDaily, updates)
        db.flush()
    result["confirmed_at"] = datetime.utcnow().isoformat() + "Z"
    stage.status, stage.result = "confirmed", result
    db.commit()
    return result


def confirm_mappings(db, stage, approved):
    """Approve only the unambiguous backfill, even if other rows block sales.

    No sales or product creation; unresolved rows never have a staged sku_id.
    """
    if not approved:
        raise HTTPException(409, "Bevestig expliciet de voorgestelde artikelcode-mapping.")
    sales_lock(db, stage.shop_id)
    stage = db.query(models.DrsDailyStage).filter_by(id=stage.id).with_for_update().populate_existing().one()
    if stage.mapping_result:
        return stage.mapping_result
    if stage.expires_at < datetime.utcnow():
        raise HTTPException(410, "Preview verlopen; maak een nieuwe preview.")
    index = IdentityIndex(db, stage.shop_id)
    count = 0
    for chunk in db.query(models.DrsDailyChunk).filter_by(stage_id=stage.id).yield_per(1):
        for row in chunk.rows:
            if not row.get("sku_id"):
                continue
            try:
                sku, proposed = index.resolve(row["code"], row["ean"])
            except ValueError as exc:
                raise HTTPException(409, f"Mapping gewijzigd: {exc}")
            if sku != row["sku_id"]:
                raise HTTPException(409, "Mapping gewijzigd sinds preview.")
            if proposed:
                discriminator = alias_discriminator(row["ean"])
                db.add(models.ProductSourceAlias(shop_id=stage.shop_id, source="DRS",
                       code=row["code"], discriminator=discriminator, sku_id=sku))
                index.aliases[(row["code"], discriminator)].add(sku)
                count += 1
            if not index.products[sku].business_sku:
                index.products[sku].business_sku = row["code"]
        db.flush()
    stage.mapping_result = {"stage_id": stage.id, "mapped": count, "sales_written": 0,
                            "unresolved": stage.preview["error_count"]}
    db.commit()
    return stage.mapping_result


def repairable(stage):
    if stage.expires_at < datetime.utcnow():
        raise HTTPException(410, "Preview verlopen; upload het volledige bestand opnieuw.")
    if stage.status == "confirmed":
        raise HTTPException(409, "Verkoop al verwerkt; upload opnieuw voor een nieuwe preview.")
    if stage.preview.get("metadata_version") != 2:
        raise HTTPException(409, "Oude preview bevat niet alle bronmetadata/fouten; upload het volledige bestand opnieuw.")


def staged_rows(db, stage):
    return [copy.deepcopy(row)
            for chunk in db.query(models.DrsDailyChunk).filter_by(stage_id=stage.id).order_by(models.DrsDailyChunk.id)
            for row in chunk.rows]


def resolution_plan(db, stage, actions):
    repairable(stage)
    if not actions or len(actions) > 1000:
        raise HTTPException(422, "Selecteer 1 tot 1000 expliciete acties per herstelvoorstel.")
    index = IdentityIndex(db, stage.shop_id)
    rows = {r["row"]: r for r in staged_rows(db, stage)}
    unresolved = {e.get("row") for e in stage.preview["errors"] if e.get("kind") == "identity"}
    current_targets = {}
    for staged in rows.values():
        try:
            target, _ = index.resolve(staged["code"], staged["ean"])
            current_targets.setdefault(target, set()).add(staged["row"])
        except ValueError:
            pass  # Still unresolved; never silently discarded by revalidation.
    used_rows, used_targets, used_keys, plans = set(), set(), set(), []
    for action in actions:
        rownum = action.get("row")
        if rownum in used_rows or rownum not in unresolved or rownum not in rows:
            raise HTTPException(409, "Dubbele actie of regel is geen herstelbare koppelfout.")
        used_rows.add(rownum)
        row = rows[rownum]
        key = (row["code"], alias_discriminator(row["ean"]))
        if key in used_keys:
            raise HTTPException(409, "Dubbele bronidentiteit in acties.")
        used_keys.add(key)
        if index.aliases.get(key):
            raise HTTPException(409, "Deze bronidentiteit is al gekoppeld; valideer opnieuw.")
        sku = action.get("sku_id")
        if action.get("action") == "link":
            if not sku or sku not in index.products:
                raise HTTPException(409, "Doelproduct bestaat niet binnen deze winkel.")
            target = index.products[sku]
            existing_eans = ean_values((target.raw_attributes or {}).get("_ean_full") or target.ean)
            supplied = ean_values(row["ean"])
            if supplied and existing_eans and not supplied & existing_eans:
                raise HTTPException(409, "EAN-conflict: bestaande identificatie wordt niet overschreven.")
            ean_targets = set().union(*(index.eans.get(e, set()) for e in supplied))
            if ean_targets and sku not in ean_targets:
                raise HTTPException(409, "EAN is aan een ander bestaand product gekoppeld; kies dat product of corrigeer de bron.")
            # An explicit source alias may differ from the canonical business SKU,
            # but it never replaces that SKU or transfers catalog ownership.
        elif action.get("action") == "create":
            if sku:
                raise HTTPException(422, "Nieuwe technische ID wordt uitsluitend door de server gemaakt.")
            if not row["product_name"] or len(row["product_name"]) > 500 or len(row["oem"]) > 100:
                raise HTTPException(409, "ProductName/OEM ontbreekt of is te lang; corrigeer en upload opnieuw.")
            if index.codes.get(row["code"]) or any(index.eans.get(e) for e in ean_values(row["ean"])):
                raise HTTPException(409, "Bestaande productcode/EAN gevonden; kies expliciet een bestaand product.")
            if sum(r["code"] == row["code"] for r in rows.values()) != 1:
                raise HTTPException(409, "Dubbele productcode in rapport; productaanmaak geblokkeerd.")
            sku = "drs-" + str(uuid4())
        else:
            raise HTTPException(422, "Onbekende herstelactie.")
        if sku in used_targets or current_targets.get(sku, set()) - {rownum}:
            raise HTTPException(409, "Meerdere rapportregels mogen niet naar hetzelfde product wijzen.")
        used_targets.add(sku)
        plans.append({"row": rownum, "action": action["action"], "sku_id": sku,
                      "code": row["code"], "product_name": row["product_name"],
                      "ean": row["ean"], "oem": row["oem"],
                      "target": index.candidate(sku) if action["action"] == "link" else None,
                      "candidates": index.evidence(row)})
    # Prevent two new products in a batch from claiming the same EAN.
    new_eans = set()
    for plan in plans:
        if plan["action"] == "create":
            values = ean_values(plan["ean"])
            if values & new_eans:
                raise HTTPException(409, "Dubbele EAN tussen productaanmaakacties.")
            new_eans.update(values)
    return index, plans


def preview_resolution(db, user, stage, actions):
    sales_lock(db, stage.shop_id)
    index, plans = resolution_plan(db, stage, actions)
    token = str(uuid4())
    expires = min(stage.expires_at, datetime.utcnow() + timedelta(minutes=30))
    impact = {"resolution_id": token, "stage_id": stage.id, "shop_id": stage.shop_id,
              "actions": plans, "products_created": sum(p["action"] == "create" for p in plans),
              "aliases_added": len(plans), "sales_written": 0,
              "requires_revalidation": True, "expires_at": expires.isoformat()}
    db.add(models.DrsIdentityResolution(id=token, stage_id=stage.id, created_by=user.username,
                                       expires_at=expires, fingerprint=index.fingerprint(), impact=impact))
    db.commit()
    return impact


def confirm_resolution(db, user, stage, resolution_id, approved):
    if not approved:
        raise HTTPException(409, "Bevestig expliciet de getoonde impact.")
    sales_lock(db, stage.shop_id)
    stage = db.query(models.DrsDailyStage).filter_by(id=stage.id).with_for_update().populate_existing().one()
    proposal = db.query(models.DrsIdentityResolution).filter_by(
        id=resolution_id, stage_id=stage.id).with_for_update().populate_existing().first()
    if not proposal or proposal.created_by != user.username:
        raise HTTPException(403, "Herstelvoorstel niet beschikbaar voor deze gebruiker/preview.")
    if proposal.result:
        return proposal.result
    repairable(stage)
    if proposal.expires_at < datetime.utcnow():
        raise HTTPException(410, "Herstelvoorstel verlopen; vraag een nieuwe impactpreview.")
    index = IdentityIndex(db, stage.shop_id)
    if index.fingerprint() != proposal.fingerprint:
        raise HTTPException(409, "Catalogus/mapping gewijzigd; vraag een nieuwe impactpreview.")
    actions = proposal.impact["actions"]
    for plan in actions:
        sku = plan["sku_id"]
        if plan["action"] == "create":
            ean = plan["ean"]
            db.add(models.DimProduct(sku_id=sku, shop_id=stage.shop_id,
                   business_sku=plan["code"], product_name=plan["product_name"],
                   ean=ean if len(ean) <= 50 else None, oem_code=plan["oem"] or None,
                   import_source="DRS", raw_attributes={"_ean_full": ean, "productcode": plan["code"],
                                                       "_identity_stage": stage.id}))
            db.flush()
        db.add(models.ProductSourceAlias(shop_id=stage.shop_id, source="DRS",
               code=plan["code"], discriminator=alias_discriminator(plan["ean"]), sku_id=sku))
        db.add(models.AuditLog(entity_type="drs_identity", entity_id=sku,
               action="drs_identity_" + plan["action"], performed_by=user.username,
               old_value=plan.get("target"), new_value=plan,
               details=json.dumps({"shop_id": stage.shop_id, "source": "DRS",
                                   "stage_id": stage.id, "resolution_id": proposal.id})))
    proposal.result = {"resolution_id": proposal.id, "status": "confirmed",
                       "products_created": proposal.impact["products_created"],
                       "aliases_added": len(actions), "sales_written": 0, "requires_revalidation": True}
    db.commit()
    return proposal.result


def revalidate_stage(db, user, stage):
    repairable(stage)
    sales_lock(db, stage.shop_id)
    return stage_rows(db, user, stage.shop_id, staged_rows(db, stage),
                      stage.preview["dates"], copy.deepcopy(stage.preview.get("parse_errors", [])),
                      stage.preview["totals"], previous=stage)