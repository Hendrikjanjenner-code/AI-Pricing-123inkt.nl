"""Reference-only cost source evaluation. Never reads or updates live products."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile

from pydantic import BaseModel, ConfigDict, Field, model_validator, StrictStr, StrictInt, StrictFloat
from fastapi import HTTPException

SOURCES = ["Stockprice", "ABCLev", "AB", "123inkt.nl Marktprijs"]


class Rules(BaseModel):
    model_config = ConfigDict(extra="forbid")
    print: list[str] = Field(default_factory=lambda: list(SOURCES))
    other: list[str] = Field(default_factory=lambda: [SOURCES[i] for i in (0, 2, 1, 3)])
    custom_source: str | None = Field(default=None, min_length=1, max_length=80)
    categories: dict[str, list[str]] = Field(default_factory=dict)

    @model_validator(mode="after")
    def valid_orders(self):
        if self.custom_source and (self.custom_source.strip() != self.custom_source or
                                   self.custom_source in [*SOURCES, "Productcode", "Category"]):
            raise ValueError("Custom source must be a distinct, nonblank name")
        sources = SOURCES + ([self.custom_source] if self.custom_source else [])
        if len(self.categories) > 50 or any(not k or len(k) > 100 for k in self.categories):
            raise ValueError("At most 50 named category overrides (100 characters)")
        for order in [self.print, self.other, *self.categories.values()]:
            if len(order) != len(sources) or set(order) != set(sources):
                raise ValueError("Each order must contain every source exactly once")
        return self


class Scenario(BaseModel):
    model_config = ConfigDict(extra="forbid")
    row_id: str = Field(min_length=1, max_length=100)
    category: str = Field(max_length=100)
    values: dict[str, StrictStr | StrictInt | StrictFloat | None] = Field(max_length=6)

    @model_validator(mode="after")
    def bound_source_names(self):
        if any(len(key) > 80 for key in self.values):
            raise ValueError("Source names may not exceed 80 characters")
        return self


class Draft(BaseModel):
    model_config = ConfigDict(extra="forbid")
    base_revision: int = Field(ge=0)
    rules: Rules


class PreviewInput(Draft):
    rows: list[Scenario] = Field(min_length=1, max_length=500)


class Activation(BaseModel):
    model_config = ConfigDict(extra="forbid")
    preview_token: str = Field(min_length=1, max_length=64)
    base_revision: int = Field(ge=0)


def order_for(rules, category):
    return rules.categories.get(category, rules.print if category.casefold() == "print" else rules.other)


def normalize_values(row, sources, required=None):
    required = set(sources if required is None else required)
    result, errors = {}, []
    if set(row.values) - set(sources):
        errors.append("Unknown source(s): " + ", ".join(sorted(set(row.values) - set(sources))))
    for source in sources:
        raw = row.values.get(source)
        if raw is None or raw == "":
            result[source] = None
            if source in required:
                errors.append(f"{source}: missing; supply explicit zero if unavailable")
            continue
        try:
            if isinstance(raw, bool) or len(str(raw)) > 80:
                raise ValueError()
            value = Decimal(str(raw))
            if (not value.is_finite() or value < 0 or value > Decimal("1000000000000")
                    or not -12 <= value.as_tuple().exponent <= 12):
                raise ValueError()
            result[source] = format(value, "f")
        except (ValueError, InvalidOperation):
            result[source] = None
            if source in required:
                errors.append(f"{source}: invalid nonnegative finite decimal (max 1e12, 12 decimal places)")
    return result, errors


def select(rules, category, values):
    trace = []
    order = order_for(rules, category)
    for i, source in enumerate(order):
        value = values.get(source)
        final = i == len(order) - 1
        if value is None:
            trace.append(dict(priority=i + 1, source=source, value=None, status="missing_or_invalid"))
            return dict(source=None, value=None, reason="input_error", fallback_warning=False, trace=trace)
        if Decimal(value) == 0 and not final:
            trace.append(dict(priority=i + 1, source=source, value=value, status="zero_skipped"))
            continue
        reason = "final_fallback" if final else "first_nonzero"
        trace.append(dict(priority=i + 1, source=source, value=value, status=reason))
        return dict(source=source, value=value, reason=reason, fallback_warning=final or source == SOURCES[3], trace=trace)


def evaluate(rows, current, proposed):
    current_sources = SOURCES + ([current.custom_source] if current.custom_source else [])
    proposed_sources = SOURCES + ([proposed.custom_source] if proposed.custom_source else [])
    sources = list(dict.fromkeys(current_sources + proposed_sources))
    results = []
    for row in rows:
        values, errors = normalize_values(row, sources, required=proposed_sources)
        warnings = []
        current_selection = select(current, row.category, values)
        incomplete = [source for source in current_sources if values[source] is None]
        if incomplete:
            warnings.append("Current comparison unavailable: missing or invalid " + ", ".join(incomplete))
            current_selection = dict(source=None, value=None, reason="current_input_incomplete",
                                     fallback_warning=True, trace=[
                                         dict(priority=i + 1, source=source, value=values[source],
                                              status="missing_or_invalid" if values[source] is None else "comparison_unavailable")
                                         for i, source in enumerate(order_for(current, row.category))])
        results.append(dict(row_id=row.row_id, category=row.category, values=values,
                            current=current_selection, proposed=select(proposed, row.category, values),
                            errors=errors, warnings=warnings))
    return results


def workbook_rows(data, rules):
    import openpyxl
    if len(data) > 2 * 1024 * 1024:
        raise HTTPException(413, "Workbook exceeds 2 MB")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 12 * 1024 * 1024 or len(archive.infolist()) > 200:
                raise HTTPException(413, "Expanded workbook too large")
        book = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        sheet = book.worksheets[0]
        if sheet.max_row > 1000 or sheet.max_column > 100:
            raise HTTPException(422, "Workbook dimensions exceed 1000 rows / 100 columns")
        iterator = sheet.iter_rows(values_only=True)
        next(iterator)
        header = list(next(iterator))
        required = ["Productcode", "Category", *SOURCES]
        if rules.custom_source:
            required.append(rules.custom_source)
        if any(header.count(name) != 1 for name in required):
            raise HTTPException(422, "Row 2 must contain each required header exactly once: " + ", ".join(required))
        indices = {name: header.index(name) for name in required}
        rows = []
        for cells in iterator:
            code = cells[indices["Productcode"]]
            if code is None:
                continue
            values = {name: cells[index] for name, index in indices.items() if name not in ("Productcode", "Category")}
            rows.append(Scenario(row_id=str(code), category=str(cells[indices["Category"]] or ""), values=values))
            if len(rows) > 500:
                raise HTTPException(422, "At most 500 scenarios")
        book.close()
        if not rows:
            raise HTTPException(422, "Workbook contains no scenario rows")
        return rows
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Invalid workbook or scenario values") from exc
