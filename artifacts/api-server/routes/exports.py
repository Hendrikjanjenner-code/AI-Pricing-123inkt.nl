"""Export endpoints — approved recommendations naar DRS: API-push, CSV/XLSX/XLS
downloads, en de bijbehorende exportgeschiedenis (exportlog)."""
import csv
import io
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy.orm import Session
import xlwt

import models
import auth as auth_utils
from connectors import drs_connector
from database import get_db

router = APIRouter(prefix="/api/exports", tags=["exports"])

EXPORT_FIELDS = [
    "sku_id", "recommended_price", "min_price", "max_price",
    "reason_code", "confidence_score", "approved_by", "date",
]


def _apply_global_filter(db: Session, recs: list,
                         hierarchy_node_ids: Optional[str],
                         brands: Optional[str]) -> list:
    """Beperk exportrijen tot de actieve globale filtercontext (taak #217).
    None = geen filter actief; lege set = niets matcht."""
    from routes.global_filters import resolve_sku_filter
    _gf = resolve_sku_filter(db, hierarchy_node_ids, brands)
    if _gf is None:
        return recs
    return [r for r in recs if r.sku_id in _gf]


def _filter_meta(db: Session, hierarchy_node_ids: Optional[str], brands: Optional[str]) -> dict:
    from routes.global_filters import filter_metadata
    return filter_metadata(db, hierarchy_node_ids, brands)


def _approved_query(db: Session, date_filter: Optional[str], only_unexported: bool = False):
    q = db.query(models.FactPricingRecommendationDaily).filter(
        models.FactPricingRecommendationDaily.approval_status == "approved",
    )
    if only_unexported:
        q = q.filter(models.FactPricingRecommendationDaily.exported == False)
    if date_filter:
        try:
            d = date.fromisoformat(date_filter)
            q = q.filter(models.FactPricingRecommendationDaily.date == d)
        except ValueError:
            pass
    return q


def _build_rows(db: Session, recs: list[models.FactPricingRecommendationDaily]) -> list[list]:
    product_map = {
        p.sku_id: p
        for p in db.query(models.DimProduct).filter(
            models.DimProduct.sku_id.in_([r.sku_id for r in recs])
        ).all()
    }
    rows = []
    for r in recs:
        product = product_map.get(r.sku_id)
        rows.append([
            r.sku_id,
            float(r.override_price or r.recommended_price) if (r.override_price or r.recommended_price) else "",
            float(product.min_price) if product and product.min_price else "",
            float(product.max_price) if product and product.max_price else "",
            r.reason_code or "",
            round(float(r.confidence_score), 3) if r.confidence_score else "",
            r.approved_by or "",
            r.date.isoformat() if r.date else "",
        ])
    return rows


def _mark_processed(db: Session, recs: list[models.FactPricingRecommendationDaily], performed_by: str, action: str) -> None:
    """Zet recommendations op exported=True + processing_status='verwerkt' en log het in audit_log."""
    now = datetime.utcnow()
    for r in recs:
        r.exported = True
        r.processing_status = "verwerkt"
        r.processed_at = now
        db.add(models.AuditLog(
            entity_type="recommendation",
            entity_id=str(r.id),
            action=action,
            performed_by=performed_by,
            details=f"sku={r.sku_id}",
        ))


def _log_export(
    db: Session,
    export_type: str,
    recs: list[models.FactPricingRecommendationDaily],
    performed_by: str,
    status: str,
    error_details: Optional[str] = None,
) -> models.ExportLog:
    entry = models.ExportLog(
        export_type=export_type,
        record_count=len(recs),
        recommendation_ids=[r.id for r in recs],
        performed_by=performed_by,
        status=status,
        error_details=error_details,
    )
    db.add(entry)
    db.commit()
    db.refresh(entry)
    return entry


def _export_log_out(e: models.ExportLog) -> dict:
    return {
        "id": e.id,
        "export_type": e.export_type,
        "record_count": e.record_count,
        "recommendation_ids": e.recommendation_ids or [],
        "performed_by": e.performed_by,
        "timestamp": e.timestamp.isoformat() if e.timestamp else None,
        "status": e.status,
        "error_details": e.error_details,
    }


@router.get("/recommendations")
def export_recommendations(
    date_filter: Optional[str] = Query(None, alias="date"),
    hierarchy_node_ids: Optional[str] = Query(None),
    brands: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth_utils.require_role("admin", "pricing_manager")),
):
    """
    Returns JSON metadata about the export.
    The actual CSV/XLSX/XLS downloads and the DRS API-push are separate endpoints below.
    """
    recs = _apply_global_filter(
        db, _approved_query(db, date_filter, only_unexported=True).all(),
        hierarchy_node_ids, brands)
    count = len(recs)
    export_date = date_filter or datetime.utcnow().date().isoformat()
    filename = f"prijsexport_{export_date}.csv"

    extra = ""
    if hierarchy_node_ids:
        extra += f"&hierarchy_node_ids={hierarchy_node_ids}"
    if brands:
        extra += f"&brands={brands}"
    return {
        "exported_count": count,
        "date": export_date,
        "download_url": f"/api/exports/recommendations/download?date={export_date}{extra}",
        "filename": filename,
        "filter_context": _filter_meta(db, hierarchy_node_ids, brands),
    }


@router.get("/recommendations/download")
def download_recommendations_csv(
    date_filter: Optional[str] = Query(None, alias="date"),
    hierarchy_node_ids: Optional[str] = Query(None),
    brands: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth_utils.require_role("admin", "pricing_manager")),
):
    """Stream the approved recommendations as a CSV file."""
    recs = _apply_global_filter(db, _approved_query(db, date_filter).all(),
                                hierarchy_node_ids, brands)
    rows = _build_rows(db, recs)

    output = io.StringIO()
    writer = csv.writer(output)
    meta = _filter_meta(db, hierarchy_node_ids, brands)
    if meta["filters_actief"]:
        # Metadata over de actieve filtercontext als commentaarregel vóór de
        # header, zodat traceerbaar is met welke selectie is geëxporteerd.
        writer.writerow([
            "# filtercontext",
            "; ".join(meta["hierarchie_selectie"]) or "-",
            "; ".join(meta["merken"]) or "-",
        ])
    writer.writerow(EXPORT_FIELDS)
    for row in rows:
        writer.writerow(row)

    _mark_processed(db, recs, current_user.username, "exported (csv)")
    _log_export(db, "csv", recs, current_user.username, "success")
    db.commit()

    output.seek(0)
    export_date = date_filter or datetime.utcnow().date().isoformat()
    filename = f"prijsexport_{export_date}.csv"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/recommendations/download/xlsx")
def download_recommendations_xlsx(
    date_filter: Optional[str] = Query(None, alias="date"),
    hierarchy_node_ids: Optional[str] = Query(None),
    brands: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth_utils.require_role("admin", "pricing_manager")),
):
    """Stream the approved recommendations as an XLSX file."""
    recs = _apply_global_filter(db, _approved_query(db, date_filter).all(),
                                hierarchy_node_ids, brands)
    rows = _build_rows(db, recs)

    wb = Workbook()
    ws = wb.active
    ws.title = "Prijsexport"
    ws.append(EXPORT_FIELDS)
    for row in rows:
        ws.append(row)

    # Metadata-tab met de actieve filtercontext (taak #217)
    meta = _filter_meta(db, hierarchy_node_ids, brands)
    if meta["filters_actief"]:
        ms = wb.create_sheet("Filters")
        ms.append(["Filtercontext van deze export"])
        ms.append(["Hiërarchie-selectie", "; ".join(meta["hierarchie_selectie"]) or "-"])
        ms.append(["Merken", "; ".join(meta["merken"]) or "-"])
        ms.append(["Geëxporteerd op", datetime.utcnow().isoformat()])
        ms.append(["Aantal rijen", len(rows)])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    _mark_processed(db, recs, current_user.username, "exported (xlsx)")
    _log_export(db, "xlsx", recs, current_user.username, "success")
    db.commit()

    export_date = date_filter or datetime.utcnow().date().isoformat()
    filename = f"prijsexport_{export_date}.xlsx"
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/recommendations/download/xls")
def download_recommendations_xls(
    date_filter: Optional[str] = Query(None, alias="date"),
    hierarchy_node_ids: Optional[str] = Query(None),
    brands: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth_utils.require_role("admin", "pricing_manager")),
):
    """Stream the approved recommendations as a legacy XLS (BIFF8) file."""
    recs = _apply_global_filter(db, _approved_query(db, date_filter).all(),
                                hierarchy_node_ids, brands)
    rows = _build_rows(db, recs)

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Prijsexport")
    for col, header in enumerate(EXPORT_FIELDS):
        ws.write(0, col, header)
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            ws.write(row_idx, col_idx, value)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    _mark_processed(db, recs, current_user.username, "exported (xls)")
    _log_export(db, "xls", recs, current_user.username, "success")
    db.commit()

    export_date = date_filter or datetime.utcnow().date().isoformat()
    filename = f"prijsexport_{export_date}.xls"
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="application/vnd.ms-excel",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.post("/recommendations/push-drs")
def push_recommendations_to_drs(
    date_filter: Optional[str] = Query(None, alias="date"),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth_utils.require_role("admin", "pricing_manager")),
):
    """
    Push approved, not-yet-exported recommendations directly to DRS via the API.
    On failure, no recommendation state changes and no data is lost — the failed
    attempt is still recorded in the exportlog with the error details.
    """
    recs = _approved_query(db, date_filter, only_unexported=True).all()
    if not recs:
        raise HTTPException(status_code=400, detail="Geen goedgekeurde, nog niet geëxporteerde adviezen gevonden.")

    rows = _build_rows(db, recs)
    payload = [dict(zip(EXPORT_FIELDS, row)) for row in rows]

    try:
        result = drs_connector.push_price_updates(payload)
    except drs_connector.ConnectorError as exc:
        _log_export(db, "api_push", recs, current_user.username, "failed", error_details=str(exc))
        raise HTTPException(status_code=502, detail=f"DRS-push mislukt: {exc}")

    _mark_processed(db, recs, current_user.username, "exported (api_push)")
    log_entry = _log_export(db, "api_push", recs, current_user.username, "success")
    db.commit()

    return {
        "export_id": log_entry.id,
        "pushed_count": len(recs),
        "accepted_count": result.get("accepted_count"),
        "status": "success",
    }


@router.get("/history")
def get_export_history(
    limit: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth_utils.require_role("admin", "pricing_manager")),
):
    """Reusable export history endpoint (e.g. for a Data Integratie dashboard)."""
    entries = (
        db.query(models.ExportLog)
        .order_by(models.ExportLog.timestamp.desc())
        .limit(limit)
        .all()
    )
    return {"items": [_export_log_out(e) for e in entries], "total": db.query(models.ExportLog).count()}
