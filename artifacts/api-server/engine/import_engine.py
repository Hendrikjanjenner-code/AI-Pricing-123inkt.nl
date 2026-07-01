"""
Data Integration Engine — centrale module voor bestandsherkenning, validatie
en kolommapping, gedeeld door alle importbronnen (PriceShape, DRS, handmatige
CSV-uploads, Paper_Copy).

Dit vervangt losse per-bron kopieën van dezelfde logica met één herbruikbare
set functies:
  - decode_csv_bytes / load_xls_sheet / load_xlsx_sheet / file_to_rows
    → bestandsherkenning voor CSV, XLSX en (legacy) XLS.
  - safe_float / safe_int / safe_date / truncate / get_alias
  - validate_ean
  - find_duplicates / build_preview
  - IMPORT_MODES / STATUS_LIFECYCLE — gedeelde vocabulaire voor versiebeheer
    en de importstatus-levenscyclus, gebruikt door elke bron.

Bronnen roepen deze module aan in plaats van elk hun eigen kopie van
parsing/validatie-logica te onderhouden.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime, date as _date
from typing import Optional

from fastapi import HTTPException

# ── Gedeelde vocabulaire: import-modes en status-levenscyclus ───────────────

# Versiebeheer-keuze die de gebruiker vóór elke import expliciet maakt.
IMPORT_MODES = ("aanvullen", "updaten", "vervangen")
DEFAULT_IMPORT_MODE = "updaten"

# Volledige statuslevenscyclus voor import_log / ImportJob, gedeeld door
# elke bron (PriceShape, DRS, handmatig).
STATUS_NIEUW = "nieuw"
STATUS_CONCEPT = "concept"
STATUS_GEVALIDEERD = "gevalideerd"
STATUS_GEIMPORTEERD = "geimporteerd"
STATUS_BIJGEWERKT = "bijgewerkt"
STATUS_OPGESLAGEN = "opgeslagen"
STATUS_FOUT = "fout"

STATUS_LIFECYCLE = (
    STATUS_NIEUW,
    STATUS_CONCEPT,
    STATUS_GEVALIDEERD,
    STATUS_GEIMPORTEERD,
    STATUS_BIJGEWERKT,
    STATUS_OPGESLAGEN,
    STATUS_FOUT,
)

PREVIEW_ROW_LIMIT = 25


# ── Bestandsdecodering ───────────────────────────────────────────────────────

def decode_csv_bytes(content: bytes) -> str:
    """
    Decode raw CSV bytes to text, robust to the encodings end users actually
    produce: UTF-8 (with/without BOM), Windows-1252/CP1252 (Excel default on
    Western-European Windows), classic Mac Roman, and Latin-1.

    Strategy:
      1. Prefer strict UTF-8 (handles BOM) — the modern default.
      2. Otherwise use charset-normalizer to detect the legacy code page so
         that special chars (€, smart quotes, accented letters) survive.
      3. Fall back to CP1252, then Latin-1 (never raises — every byte maps).
    """
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass

    try:
        from charset_normalizer import from_bytes

        match = from_bytes(
            content,
            cp_isolation=["utf_8", "cp1252", "mac_roman", "latin_1"],
        ).best()
        if match is not None and match.encoding:
            try:
                return content.decode(match.encoding)
            except (UnicodeDecodeError, LookupError):
                pass
    except ImportError:
        pass

    try:
        return content.decode("cp1252")
    except UnicodeDecodeError:
        return content.decode("latin-1")


def xlsx_worksheet_to_rows(ws) -> list[dict]:
    """
    Convert a single openpyxl worksheet object into row dicts (lower-cased
    header keys). Shared by both the generic single-sheet loader below and
    multi-tab sources (DRS) that need to pick a specific sheet by name.
    """
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(c).strip().lower() if c is not None else "" for c in next(rows_iter, [])]
    rows: list[dict] = []
    for row in rows_iter:
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        row_dict: dict = {}
        for h, v in zip(headers, row):
            row_dict[h] = "" if v is None else str(v).strip()
        rows.append(row_dict)
    return rows


def open_xlsx_workbook(content: bytes):
    """Open an XLSX workbook (read-only, values only). Raises HTTPException on failure."""
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl is niet geïnstalleerd; XLSX niet ondersteund")
    try:
        return openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"XLSX bestand kon niet worden geopend: {e}")


def load_xlsx_sheet(content: bytes, sheet_index: int = 0, warnings: list | None = None) -> list[dict]:
    """Load one worksheet (default: the first) of an XLSX workbook into row dicts."""
    wb = open_xlsx_workbook(content)
    if len(wb.worksheets) > 1 and warnings is not None and sheet_index == 0:
        sheet_names = ", ".join(ws.title for ws in wb.worksheets)
        warnings.append(
            f"Alleen het eerste werkblad wordt ingelezen (werkbladen gevonden: {sheet_names})"
        )
    ws = wb.worksheets[sheet_index]
    rows = xlsx_worksheet_to_rows(ws)
    wb.close()
    return rows


def open_xls_workbook(content: bytes):
    """Open a legacy .xls (BIFF) workbook via xlrd. Raises HTTPException on failure."""
    try:
        import xlrd
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="xlrd is niet geïnstalleerd; legacy XLS-bestanden niet ondersteund",
        )
    try:
        return xlrd.open_workbook(file_contents=content)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"XLS bestand kon niet worden geopend: {e}")


def xls_sheet_to_rows(wb, sheet) -> list[dict]:
    """Convert a single xlrd sheet object into row dicts (lower-cased header keys)."""
    import xlrd

    if sheet.nrows == 0:
        return []

    def _cell_to_str(cell) -> str:
        if cell.ctype == xlrd.XL_CELL_EMPTY or cell.value is None:
            return ""
        if cell.ctype == xlrd.XL_CELL_DATE:
            try:
                dt = xlrd.xldate_as_datetime(cell.value, wb.datemode)
                return dt.date().isoformat()
            except Exception:
                return str(cell.value)
        if cell.ctype == xlrd.XL_CELL_NUMBER:
            # Render whole numbers without a trailing ".0" (Excel stores all
            # numbers as float internally).
            val = cell.value
            if isinstance(val, float) and val.is_integer():
                return str(int(val))
            return str(val)
        return str(cell.value).strip()

    headers = [_cell_to_str(sheet.cell(0, c)).lower() for c in range(sheet.ncols)]
    rows: list[dict] = []
    for r in range(1, sheet.nrows):
        values = [_cell_to_str(sheet.cell(r, c)) for c in range(sheet.ncols)]
        if all(v == "" for v in values):
            continue
        rows.append(dict(zip(headers, values)))
    return rows


def load_xls_sheet(content: bytes, sheet_index: int = 0, warnings: list | None = None) -> list[dict]:
    """
    Load one worksheet (default: the first) of a legacy .xls (BIFF) workbook
    into row dicts, via xlrd. Provides the same "old Excel" support as XLSX
    for every source (PriceShape, DRS, handmatig).
    """
    wb = open_xls_workbook(content)
    if wb.nsheets > 1 and warnings is not None and sheet_index == 0:
        sheet_names = ", ".join(wb.sheet_names())
        warnings.append(
            f"Alleen het eerste werkblad wordt ingelezen (werkbladen gevonden: {sheet_names})"
        )
    sheet = wb.sheet_by_index(sheet_index)
    return xls_sheet_to_rows(wb, sheet)


def file_to_rows(content: bytes, filename: str, warnings: list | None = None) -> list[dict]:
    """
    Parse uploaded file bytes into a list of row dicts. Supports .csv
    (UTF-8/CP1252/Mac-Roman/Latin-1), .xlsx and legacy .xls (first sheet only
    for spreadsheet formats). Values are always returned as strings, matching
    csv.DictReader behaviour. All import sources (PriceShape, DRS, handmatig)
    call this single entry point.
    """
    fname = (filename or "").lower()
    if fname.endswith(".xlsx"):
        return load_xlsx_sheet(content, warnings=warnings)
    if fname.endswith(".xls"):
        return load_xls_sheet(content, warnings=warnings)

    text = decode_csv_bytes(content)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    reader = csv.DictReader(io.StringIO(text))
    return [{(k or "").strip().lower(): v for k, v in row.items()} for row in reader]


# ── Waarde-parsers ───────────────────────────────────────────────────────────

def safe_float(val) -> Optional[float]:
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    try:
        return float(s.replace(",", "."))
    except (ValueError, TypeError):
        return None


def safe_int(val) -> Optional[int]:
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    try:
        return int(float(s))
    except (ValueError, TypeError):
        return None


def safe_date(val) -> Optional[_date]:
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def truncate(value: Optional[str], limit: int) -> Optional[str]:
    """Truncate string to a column's VARCHAR limit to prevent StringDataRightTruncation."""
    if value is None:
        return None
    return value[:limit] if len(value) > limit else value


def get_alias(row: dict, *keys: str) -> Optional[str]:
    """Return the first non-empty value among alias column names (case-insensitive)."""
    for k in keys:
        v = row.get(k) or row.get(k.lower()) or row.get(k.upper())
        if v is not None and str(v).strip():
            return str(v).strip()
    return None


def validate_ean(ean: str) -> bool:
    """Check EAN-13 or EAN-8 check-digit. Returns True if valid."""
    ean = (ean or "").strip().replace(" ", "")
    if not ean.isdigit():
        return False
    if len(ean) == 13:
        total = sum(int(d) * (1 if i % 2 == 0 else 3) for i, d in enumerate(ean[:12]))
        return (10 - (total % 10)) % 10 == int(ean[12])
    if len(ean) == 8:
        weights = [3, 1, 3, 1, 3, 1, 3]
        total = sum(int(d) * w for d, w in zip(ean[:7], weights))
        return (10 - (total % 10)) % 10 == int(ean[7])
    return False


# ── Validatie-primitieven (herbruikbaar door elke bron) ─────────────────────

def find_duplicates(rows: list[dict], key_fn) -> dict[str, list[int]]:
    """
    Return {key: [1-indexed row numbers]} for every key produced by key_fn
    that occurs more than once (empty keys are ignored).
    """
    seen: dict[str, list[int]] = {}
    for i, row in enumerate(rows, start=1):
        key = key_fn(row)
        if key:
            seen.setdefault(key, []).append(i)
    return {k: v for k, v in seen.items() if len(v) > 1}


def build_preview(rows: list[dict], limit: int = PREVIEW_ROW_LIMIT) -> list[dict]:
    """Return the first `limit` rows for a validation preview (default 25 rows)."""
    return rows[:limit]


def check_required_groups(row_norm: dict, groups: list[list[str]]) -> list[str]:
    """
    Given a lower-cased row dict and a list of alias groups (each group is a
    list of acceptable column-name aliases for one logical required field),
    return the canonical field names (first alias of each group) that are
    missing/empty in this row.
    """
    missing = []
    for group in groups:
        if not any((row_norm.get(alias.lower()) or "").strip() for alias in group):
            missing.append(group[0])
    return missing


def missing_columns(available_columns: set[str], groups: list[list[str]]) -> list[str]:
    """Return canonical field names whose alias group has no match in available_columns."""
    missing = []
    for group in groups:
        if not any(alias.lower() in available_columns for alias in group):
            missing.append(group[0])
    return missing
