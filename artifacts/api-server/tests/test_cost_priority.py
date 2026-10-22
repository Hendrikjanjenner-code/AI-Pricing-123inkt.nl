"""Isolated step-1 checks. Run with --noconftest: no application lifecycle."""
import os
import sys
from pathlib import Path
from types import SimpleNamespace
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from uuid import uuid4

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from services.cost_priority import Rules, Scenario, PreviewInput, Activation, SOURCES, evaluate, workbook_rows
from routes import cost_priority as api
from models import DimShop, CostPriorityVersion, CostPriorityPreview, AuditLog

WORKBOOK = Path(__file__).resolve().parents[3] / "attached_assets/PriceShapeTestABCLevLaag_1792567013080.xlsx"
USER = SimpleNamespace(id=41, username="manager", role="pricing_manager")
TABLES = [DimShop.__table__, CostPriorityVersion.__table__, CostPriorityPreview.__table__, AuditLog.__table__]


@pytest.fixture
def store(monkeypatch):
    engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})
    @event.listens_for(engine, "connect")
    def fk(conn, _):
        conn.execute("PRAGMA foreign_keys=ON")
    for table in TABLES:
        table.create(engine)
    factory = sessionmaker(bind=engine)
    with factory() as db:
        db.add_all([DimShop(id=i, shop_code=str(i), name=str(i)) for i in (1, 2)])
        db.commit()
    monkeypatch.setattr(api, "effective_permissions", lambda db, user: {"pricing.view", "pricing.propose", "pricing.approve"})
    monkeypatch.setattr(api, "user_allowed_shop_ids", lambda db, user: {1})
    yield factory
    engine.dispose()


def body(rules=None, values=None):
    return PreviewInput(base_revision=0, rules=rules or Rules(),
                        rows=[Scenario(row_id="reference", category="Office",
                                       values=values or dict(zip(SOURCES, ["0", "12", "10", "15"])))])


def test_all_workbook_formulas_and_cached_results():
    import openpyxl
    rows = workbook_rows(WORKBOOK.read_bytes(), Rules())
    formulas = openpyxl.load_workbook(WORKBOOK, data_only=False).active
    cached = openpyxl.load_workbook(WORKBOOK, data_only=True).active
    results = evaluate(rows, Rules(), Rules())
    assert len(results) == 8
    for i, result in enumerate(results, 3):
        # Every actual formula row checked, including Office tie in row 9.
        assert formulas.cell(i, 12).value == (
            f'=IF(C{i}="Print",IF(E{i}<>0,$E$2,IF(F{i}<>0,$F$2,IF(G{i}<>0,$G$2,$H$2))),'
            f'IF(E{i}<>0,$E$2,IF(G{i}<>0,$G$2,IF(F{i}<>0,$F$2,$H$2))))')
        assert formulas.cell(i, 10).value == f'=IF(L{i}=$E$2,E{i},IF(L{i}=$F$2,F{i},IF(L{i}=$G$2,G{i},H{i})))'
        assert result["proposed"]["source"] == cached.cell(i, 12).value
        assert result["proposed"]["value"] == str(cached.cell(i, 10).value)
        assert not result["errors"]
    assert results[6]["proposed"]["source"] == "AB"  # prose column K incorrectly says ABCLev


def test_zero_missing_invalid_reorder_and_custom():
    zeros = body(values=dict.fromkeys(SOURCES, "0")).rows
    result = evaluate(zeros, Rules(), Rules())[0]
    assert result["proposed"]["source"] == SOURCES[-1]
    assert result["proposed"]["value"] == "0"
    assert len(result["proposed"]["trace"]) == 4
    assert result["proposed"]["fallback_warning"]
    for bad in (None, "-1", "NaN", "Infinity", "=SUM(A1:A2)", "0e999999"):
        row = body(values={**dict.fromkeys(SOURCES, "0"), "ABCLev": bad}).rows
        assert evaluate(row, Rules(), Rules())[0]["errors"]
    reordered = Rules(other=["ABCLev", "AB", "Stockprice", SOURCES[-1]])
    result = evaluate(body().rows, Rules(), reordered)[0]
    assert result["current"]["value"] == "10" and result["proposed"]["value"] == "12"
    custom = Rules(custom_source="Manual contract", print=["Manual contract", *SOURCES],
                   other=["Manual contract", *SOURCES])
    custom_rows = body(values={**dict.fromkeys(SOURCES, "0"), "Manual contract": "2.345"}).rows
    assert evaluate(custom_rows, Rules(), custom)[0]["proposed"]["value"] == "2.345"
    with pytest.raises(ValueError):
        Rules(print=["AB"] * 4)
    with pytest.raises(ValueError):
        body(values={**dict.fromkeys(SOURCES, 0), "AB": True})


def test_activation_reload_history_stale_and_cascade(store):
    with store() as db:
        assert api.get_configuration(1, db, USER)["revision"] == 0
        assert db.query(CostPriorityVersion).count() == 0
        first = api.preview_manual(1, body(), db, USER)
        changed = Rules(other=["ABCLev", "AB", "Stockprice", SOURCES[-1]])
        second = api.preview_manual(1, body(changed), db, USER)
        with pytest.raises(HTTPException) as exc:
            api.activate(1, Activation(preview_token=first["preview_token"], base_revision=0), db, USER)
        assert exc.value.status_code == 409
        request = Activation(preview_token=second["preview_token"], base_revision=0)
        assert api.activate(1, request, db, USER)["revision"] == 1
        with pytest.raises(HTTPException) as exc:
            api.activate(1, request, db, USER)
        assert exc.value.status_code == 409
        assert api.get_configuration(1, db, USER)["rules"] == changed.model_dump()
        assert len(api.history(1, db, USER)["versions"]) == 1
        assert db.query(AuditLog).count() == 1
        db.execute(text("DELETE FROM dim_shop WHERE id=1"))
        db.commit()
        assert db.query(CostPriorityVersion).count() == db.query(CostPriorityPreview).count() == 0


def test_error_activation_and_tenant_role_permissions(store, monkeypatch):
    with store() as db:
        bad = api.preview_manual(1, body(values=dict.fromkeys(SOURCES, None)), db, USER)
        assert not bad["can_activate"]
        with pytest.raises(HTTPException) as exc:
            api.activate(1, Activation(preview_token=bad["preview_token"], base_revision=0), db, USER)
        assert exc.value.status_code == 422
        for call in (lambda: api.get_configuration(2, db, USER),
                     lambda: api.preview_manual(2, body(), db, USER),
                     lambda: api.history(2, db, USER),
                     lambda: api.activate(2, Activation(preview_token=bad["preview_token"], base_revision=0), db, USER),
                     lambda: api.preview_manual(1, body(), db, SimpleNamespace(id=9, role="viewer"))):
            with pytest.raises(HTTPException) as exc:
                call()
            assert exc.value.status_code == 403
        monkeypatch.setattr(api, "effective_permissions", lambda db, user: {"pricing.view"})
        assert not api.get_configuration(1, db, USER)["can_edit"]
        with pytest.raises(HTTPException) as exc:
            api.preview_manual(1, body(), db, USER)
        assert exc.value.status_code == 403


def test_http_manual_and_workbook(store):
    app = FastAPI()
    app.include_router(api.router)
    def database():
        with store() as db:
            yield db
    app.dependency_overrides[api.get_db] = database
    app.dependency_overrides[api.get_current_user] = lambda: USER
    with TestClient(app) as client:
        assert client.get("/api/cost-priority/1").status_code == 200
        response = client.post("/api/cost-priority/1/preview", json=body().model_dump())
        assert response.status_code == 200
        response = client.post("/api/cost-priority/1/preview-workbook",
                               data={"draft": '{"base_revision":0,"rules":' + Rules().model_dump_json() + '}'},
                               files={"file": ("reference.xlsx", WORKBOOK.read_bytes())})
        assert response.status_code == 200
        assert len(response.json()["rows"]) == 8
        assert response.json()["can_activate"]
        assert client.post("/api/cost-priority/1/activate", json={
            "preview_token": response.json()["preview_token"], "base_revision": 0}).status_code == 200


def test_cross_user_stale_revision_and_expiry(store):
    from datetime import datetime, timedelta
    other = SimpleNamespace(id=42, username="other-manager", role="pricing_manager")
    with store() as db:
        first = api.preview_manual(1, body(), db, USER)
        second = api.preview_manual(1, body(), db, other)
        with pytest.raises(HTTPException) as exc:
            api.activate(1, Activation(preview_token=first["preview_token"], base_revision=0), db, other)
        assert exc.value.status_code == 409
        api.activate(1, Activation(preview_token=first["preview_token"], base_revision=0), db, USER)
        with pytest.raises(HTTPException) as exc:
            api.activate(1, Activation(preview_token=second["preview_token"], base_revision=0), db, other)
        assert exc.value.status_code == 409
        with pytest.raises(HTTPException) as exc:
            api.preview_manual(1, body(), db, USER)
        assert exc.value.status_code == 409
        draft = body()
        draft.base_revision = 1
        fresh = api.preview_manual(1, draft, db, USER)
        db.get(CostPriorityPreview, fresh["preview_token"]).expires_at = datetime.utcnow() - timedelta(seconds=1)
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.activate(1, Activation(preview_token=fresh["preview_token"], base_revision=1), db, USER)
        assert exc.value.status_code == 409


def test_workbook_formula_inputs_rejected_and_stock_quantity_ignored():
    from io import BytesIO
    import openpyxl
    book = openpyxl.load_workbook(WORKBOOK)
    book.active["D3"] = 0
    output = BytesIO()
    book.save(output)
    result = evaluate(workbook_rows(output.getvalue(), Rules()), Rules(), Rules())[0]
    assert result["proposed"]["source"] == "Stockprice"
    book.active["E3"] = "=1+1"
    output = BytesIO()
    book.save(output)
    result = evaluate(workbook_rows(output.getvalue(), Rules()), Rules(), Rules())[0]
    assert result["errors"]
    assert result["proposed"]["source"] is None


@pytest.mark.parametrize("kind", ["manual", "workbook"])
@pytest.mark.parametrize("replacement", [None, "New contract"])
def test_remove_or_rename_active_custom_source(store, kind, replacement):
    from io import BytesIO
    import json
    import openpyxl

    old = "Old contract"
    old_rules = Rules(custom_source=old, print=[old, *SOURCES], other=[old, *SOURCES])
    new_rules = (Rules(custom_source=replacement, print=[replacement, *SOURCES],
                       other=[replacement, *SOURCES]) if replacement else Rules())
    with store() as db:
        initial = api.preview_manual(1, body(old_rules, {**dict.fromkeys(SOURCES, "10"), old: "5"}), db, USER)
        api.activate(1, Activation(preview_token=initial["preview_token"], base_revision=0), db, USER)
    app = FastAPI()
    app.include_router(api.router)
    def database():
        with store() as db:
            yield db
    app.dependency_overrides[api.get_db] = database
    app.dependency_overrides[api.get_current_user] = lambda: USER
    with TestClient(app) as client:
        draft = {"base_revision": 1, "rules": new_rules.model_dump()}
        if kind == "manual":
            values = dict.fromkeys(SOURCES, "10")
            if replacement:
                values[replacement] = "7"
            response = client.post("/api/cost-priority/1/preview", json={
                **draft, "rows": [{"row_id": "example", "category": "Print", "values": values}]})
        else:
            book = openpyxl.load_workbook(WORKBOOK)
            if replacement:
                book.active.cell(2, 13, replacement)
                for i in range(3, 11):
                    book.active.cell(i, 13, 7)
            output = BytesIO()
            book.save(output)
            response = client.post("/api/cost-priority/1/preview-workbook",
                                   data={"draft": json.dumps(draft)},
                                   files={"file": ("reference.xlsx", output.getvalue())})
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["can_activate"] and not result["errors"]
        for row in result["rows"]:
            assert row["warnings"] and not row["errors"]
            assert row["current"]["source"] is None and row["current"]["value"] is None
            assert row["current"]["reason"] == "current_input_incomplete"
            assert row["current"]["fallback_warning"]
            assert row["proposed"]["source"] is not None
        activated = client.post("/api/cost-priority/1/activate", json={
            "preview_token": result["preview_token"], "base_revision": 1})
        assert activated.status_code == 200, activated.text
        assert activated.json()["revision"] == 2
        assert activated.json()["rules"] == new_rules.model_dump()


def test_available_shops_pricing_view_only_tenant_filter(store, monkeypatch):
    app = FastAPI()
    app.include_router(api.router)
    def database():
        with store() as db:
            yield db
    app.dependency_overrides[api.get_db] = database
    app.dependency_overrides[api.get_current_user] = lambda: USER
    monkeypatch.setattr(api, "effective_permissions", lambda db, user: {"pricing.view"})
    with store() as db:
        db.add(DimShop(id=3, shop_code="inactive", name="Inactive", active=False))
        db.commit()
    with TestClient(app) as client:
        response = client.get("/api/cost-priority/shops/available")
        assert response.status_code == 200
        assert [shop["id"] for shop in response.json()["shops"]] == [1]
        assert set(response.json()["shops"][0]) == {"id", "name", "is_default", "active"}
        monkeypatch.setattr(api, "user_allowed_shop_ids", lambda db, user: set())
        assert client.get("/api/cost-priority/shops/available").json() == {"shops": []}
        monkeypatch.setattr(api, "user_allowed_shop_ids", lambda db, user: None)
        assert [shop["id"] for shop in client.get("/api/cost-priority/shops/available").json()["shops"]] == [1, 2]
        monkeypatch.setattr(api, "user_allowed_shop_ids", lambda db, user: {3, 99})
        assert client.get("/api/cost-priority/shops/available").json() == {"shops": []}
        monkeypatch.setattr(api, "effective_permissions", lambda db, user: set())
        assert client.get("/api/cost-priority/shops/available").status_code == 403


def test_optional_current_only_values_and_reserved_names():
    old = Rules(custom_source="Old", print=["Old", *SOURCES], other=["Old", *SOURCES])
    values = {**dict.fromkeys(SOURCES, "10"), "Old": "5"}
    row = Scenario(row_id="reference", category="Print", values=values)
    result = evaluate([row], old, Rules())[0]
    assert result["current"]["value"] == "5"
    assert not result["warnings"] and not result["errors"]
    row.values["Old"] = "-1"
    result = evaluate([row], old, Rules())[0]
    assert result["warnings"] and not result["errors"]
    assert result["current"]["value"] is None
    for reserved in ("Productcode", "Category"):
        with pytest.raises(ValueError):
            Rules(custom_source=reserved, print=[reserved, *SOURCES], other=[reserved, *SOURCES])


def test_postgres_concurrent_activation_isolated_schema(monkeypatch):
    """Dedicated test DB plus random schema; never uses application tables."""
    from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode
    import psycopg2
    from psycopg2.extensions import parse_dsn
    original = os.environ.get("APP_DATABASE_URL_ORIGINAL") or os.environ["DATABASE_URL"]
    parts = urlsplit(original)
    app_name = parse_dsn(original)["dbname"]
    if os.environ.get("APP_DATABASE_URL_ORIGINAL"):
        # conftest has already selected and guarded the test database.
        test_url = os.environ["DATABASE_URL"]
    else:
        test_url = os.environ.get("TEST_DATABASE_URL") or urlunsplit(parts._replace(
            path="/" + app_name + "_test",
            query=urlencode([(k, v) for k, v in parse_qsl(parts.query) if k.lower() != "dbname"])))
    assert parse_dsn(test_url)["dbname"] != app_name
    # Test DB must already exist; this test never creates a database or touches app data.
    admin = create_engine(test_url)
    schema = "cost_priority_" + uuid4().hex
    with admin.begin() as conn:
        conn.execute(text(f'CREATE SCHEMA "{schema}"'))
    engine = create_engine(test_url, connect_args={"options": f"-c search_path={schema} -c statement_timeout=10000"})
    try:
        for table in TABLES:
            table.create(engine)
        factory = sessionmaker(bind=engine)
        monkeypatch.setattr(api, "effective_permissions", lambda db, user: {"pricing.view", "pricing.propose", "pricing.approve"})
        monkeypatch.setattr(api, "user_allowed_shop_ids", lambda db, user: {1})
        with factory() as db:
            db.add(DimShop(id=1, shop_code="isolated", name="isolated"))
            db.commit()
            p = api.preview_manual(1, body(), db, USER)
        barrier = Barrier(2)
        def worker():
            with factory() as db:
                barrier.wait(timeout=10)
                try:
                    return api.activate(1, Activation(preview_token=p["preview_token"], base_revision=0), db, USER)["revision"]
                except HTTPException as exc:
                    return exc.status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            assert sorted(pool.map(lambda _: worker(), range(2))) == [1, 409]
        with factory() as db:
            assert db.query(CostPriorityVersion).count() == 1
            assert db.query(AuditLog).count() == 1
            db.execute(text("DELETE FROM dim_shop WHERE id=1"))
            db.commit()
            assert db.query(CostPriorityVersion).count() == db.query(CostPriorityPreview).count() == 0
    finally:
        engine.dispose()
        with admin.begin() as conn:
            conn.execute(text(f'DROP SCHEMA "{schema}" CASCADE'))
        admin.dispose()
