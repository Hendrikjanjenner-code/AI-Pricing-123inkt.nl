"""Daily DRS invariants. conftest routes all writes to the isolated test DB."""
import io
import os
import sys
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from uuid import uuid4
from concurrent.futures import ThreadPoolExecutor
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
import openpyxl
import pytest
from fastapi import HTTPException
from sqlalchemy.exc import IntegrityError
from database import SessionLocal
import models
from services.drs_daily import parse_workbook, create_preview, confirm_stage, confirm_mappings
from services.product_identity import IdentityIndex, authorize_shop


def workbook(code="0007", ean="001234", values=None, days=("30/08", "31/08"), wrong_total=False):
    values = values or [(2, "24.2", "20"), (-1, "-12.1", "-10")]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sales"
    ws.append([None] * 4 + [x for d in days for x in (d, None, None)] + ["Totaal"])
    ws.append(["Productcode", "ProductName", "EAN", "OEM"] +
              ["Quantity", "Revenue Incl VAT", "Revenue Excl. VAT"] * (len(days) + 1))
    totals = [sum(Decimal(str(v[i])) for v in values) for i in range(3)]
    if wrong_total:
        totals[0] += 1
    ws.append([code, "Product", ean, "NONUNIQUE"] +
              [float(x) for v in values for x in v] + [float(x) for x in totals])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


@pytest.fixture
def db():
    db = SessionLocal()
    suffix = uuid4().hex[:12]
    shops = [models.DimShop(shop_code=f"daily-{suffix}-{i}", name="Daily test") for i in range(2)]
    db.add_all(shops)
    db.flush()
    products = [models.DimProduct(sku_id=f"daily-{suffix}-{i}", shop_id=s.id,
                product_name="Identity fixture", ean="001234", product_group="0007",
                import_source="priceshape", cost_price=100) for i, s in enumerate(shops)]
    db.add_all(products)
    db.commit()
    db.test_shops, db.test_products = shops, products
    yield db
    db.rollback()
    from sqlalchemy import text
    db.execute(text("SET LOCAL statement_timeout = '60000'"))  # isolated large-fixture cleanup only
    ids, skus = [s.id for s in shops], [p.sku_id for p in products]
    stage_ids = [s.id for s in db.query(models.DrsDailyStage).filter(models.DrsDailyStage.shop_id.in_(ids))]
    db.query(models.DrsIdentityResolution).filter(models.DrsIdentityResolution.stage_id.in_(stage_ids)).delete(synchronize_session=False)
    owned_skus = [s for s, in db.query(models.DimProduct.sku_id).filter(models.DimProduct.shop_id.in_(ids))]
    db.query(models.AuditLog).filter(models.AuditLog.entity_type == "drs_identity", models.AuditLog.entity_id.in_(owned_skus)).delete(synchronize_session=False)
    db.query(models.DrsDailyChunk).filter(models.DrsDailyChunk.stage_id.in_(stage_ids)).delete(synchronize_session=False)
    db.query(models.DrsDailyStage).filter(models.DrsDailyStage.shop_id.in_(ids)).delete(synchronize_session=False)
    db.query(models.SalesDailyClaim).filter(models.SalesDailyClaim.shop_id.in_(ids)).delete(synchronize_session=False)
    db.query(models.FactSalesDaily).filter(models.FactSalesDaily.shop_id.in_(ids)).delete(synchronize_session=False)
    db.query(models.ProductSourceAlias).filter(models.ProductSourceAlias.shop_id.in_(ids)).delete(synchronize_session=False)
    db.query(models.DimProduct).filter(models.DimProduct.shop_id.in_(ids)).delete(synchronize_session=False)
    db.query(models.DimShop).filter(models.DimShop.id.in_(ids)).delete(synchronize_session=False)
    db.commit()
    db.close()


def preview(db, values=None, shop=0, **kwargs):
    return create_preview(db, SimpleNamespace(username="test"), db.test_shops[shop].id,
                          workbook(values=values, **kwargs), 2026, date(2026, 8, 30), date(2026, 8, 31), {})


def confirm(db, p):
    return confirm_stage(db, db.get(models.DrsDailyStage, p["stage_id"]), True)


def test_repeat_overlap_correction_zero_and_signed_values(db):
    p = preview(db)
    assert p["status"] == "ready" and p["totals"]["quantity"] == 1
    assert p["mapping_proposal_count"] == 1
    r = confirm(db, p)
    assert r["inserted"] == 2
    from datetime import datetime, timezone
    timestamp = datetime.fromisoformat(r["confirmed_at"].replace("Z", "+00:00"))
    assert 0 <= (datetime.now(timezone.utc) - timestamp).total_seconds() < 60
    db.expire_all()
    assert db.get(models.DrsDailyStage, p["stage_id"]).result["confirmed_at"] == r["confirmed_at"]
    assert confirm(db, p) == r
    assert confirm(db, preview(db))["unchanged"] == 2
    correction = preview(db, [(0, 0, 0), (-2, "-24.2", "-20")])
    assert confirm(db, correction)["updated"] == 2
    facts = db.query(models.FactSalesDaily).filter_by(shop_id=db.test_shops[0].id).order_by(models.FactSalesDaily.date).all()
    assert [f.units_sold for f in facts] == [0, -2]
    assert facts[1].revenue == -20 and facts[1].revenue_incl_vat == Decimal("-24.2")
    assert all(f.cost is None and f.margin is None for f in facts)
    assert db.test_products[0].sku_id == facts[0].sku_id
    assert db.test_products[0].business_sku == "0007"
    assert db.test_products[1].business_sku is None


def test_legacy_overlap_blocks_even_zero_and_is_rechecked(db):
    p = preview(db, [(0, 0, 0), (0, 0, 0)])
    db.add(models.FactSalesDaily(shop_id=db.test_shops[0].id, sku_id=db.test_products[0].sku_id,
                                date=date(2026, 8, 30), units_sold=1, revenue=10))
    db.commit()
    assert preview(db)["status"] == "blocked"
    with pytest.raises(HTTPException, match="409"):
        confirm(db, p)
    db.rollback()
    assert db.query(models.SalesDailyClaim).filter_by(shop_id=db.test_shops[0].id).count() == 0


def test_mapping_ambiguity_conflict_and_shop_isolation(db):
    shop = db.test_shops[0].id
    duplicate = models.DimProduct(sku_id=f"amb-{uuid4().hex}", shop_id=shop,
                                 product_name="Ambiguous", product_group="0007", ean="009999")
    db.add(duplicate)
    db.commit()
    index = IdentityIndex(db, shop)
    assert index.resolve("0007", "001234")[0] == db.test_products[0].sku_id
    with pytest.raises(ValueError, match="Ambigue"):
        index.resolve("0007")
    with pytest.raises(ValueError, match="EAN-conflict"):
        index.resolve("0007", "008888")
    with pytest.raises(ValueError, match="Onbekende"):
        index.resolve("007", "001234")
    assert preview(db, ean="")["status"] == "blocked"
    assert preview(db, shop=1)["status"] == "ready"
    assert confirm(db, preview(db, shop=1))["inserted"] == 2
    assert db.query(models.FactSalesDaily).filter_by(shop_id=shop).count() == 0


def test_expired_stage_and_mapping_approval_required(db):
    from datetime import datetime, timedelta
    p = preview(db)
    stage = db.get(models.DrsDailyStage, p["stage_id"])
    with pytest.raises(HTTPException) as exc:
        confirm_stage(db, stage, False)
    assert exc.value.status_code == 409
    db.rollback()
    stage.expires_at = datetime.utcnow() - timedelta(seconds=1)
    db.commit()
    with pytest.raises(HTTPException) as exc:
        confirm(db, p)
    assert exc.value.status_code == 410
    db.rollback()


def test_shop_authorization_and_alias_mismatch(db, monkeypatch):
    import services.product_identity as identity
    monkeypatch.setattr(identity, "user_allowed_shop_ids", lambda *_: {db.test_shops[0].id})
    with pytest.raises(HTTPException) as exc:
        authorize_shop(db, SimpleNamespace(id=1), db.test_shops[1].id)
    assert exc.value.status_code == 403
    p = preview(db)
    db.test_products[0].ean = "009999"
    db.commit()
    with pytest.raises(HTTPException) as exc:
        confirm(db, p)
    assert exc.value.status_code == 409
    db.rollback()


def test_concurrent_confirm_and_database_claim_uniqueness(db):
    p = preview(db)
    stage_id = p["stage_id"]
    def worker():
        with SessionLocal() as session:
            return confirm_stage(session, session.get(models.DrsDailyStage, stage_id), True)
    with ThreadPoolExecutor(max_workers=2) as pool:
        a, b = list(pool.map(lambda _: worker(), range(2)))
    assert a == b and a["inserted"] == 2
    claim = models.SalesDailyClaim(shop_id=db.test_shops[0].id, sku_id=db.test_products[0].sku_id,
                                  date=date(2026, 8, 30), source="OTHER")
    db.add(claim)
    with pytest.raises(IntegrityError):
        db.commit()
    db.rollback()


def test_long_competing_confirmation_is_explicit_conflict_not_sql_timeout(db):
    from services.product_identity import sales_lock
    shop = db.test_shops[0].id
    sales_lock(db, shop)
    with SessionLocal() as contender:
        with pytest.raises(HTTPException) as exc:
            sales_lock(contender, shop)
        assert exc.value.status_code == 409
        contender.rollback()
    db.rollback()


def test_parser_year_boundary_totals_mapping_and_leading_zeroes():
    rows, days, errors, totals = parse_workbook(workbook(days=("31/12", "01/01")),
                                               2026, date(2026, 12, 31), date(2027, 1, 1))
    assert days == ["2026-12-31", "2027-01-01"] and not errors
    assert rows[0]["code"] == "0007" and rows[0]["ean"] == "001234"
    assert totals["revenue_excl_vat"] == "10.0000"
    assert parse_workbook(workbook(wrong_total=True), 2026, date(2026, 8, 30), date(2026, 8, 31))[2]
    with pytest.raises(ValueError, match="periode"):
        parse_workbook(workbook(), 2025, date(2026, 8, 30), date(2026, 8, 31))
    with pytest.raises(ValueError, match="kolomkoppeling"):
        parse_workbook(workbook(), 2026, date(2026, 8, 30), date(2026, 8, 31), {"unknown": 1})


def test_multi_ean_is_set_not_oem_or_arbitrary_first_match(db):
    index = IdentityIndex(db, db.test_shops[0].id)
    assert index.resolve("0007", "999999, 001234")[0] == db.test_products[0].sku_id
    assert preview(db, ean="999999, 001234")["status"] == "ready"
    p = preview(db, ean=", ".join(["001234"] + [str(i) for i in range(30)]))
    assert confirm(db, p)["inserted"] == 2
    assert IdentityIndex(db, db.test_shops[0].id).resolve("0007", "999999, 001234")[0] == db.test_products[0].sku_id


def test_safe_backfill_can_approve_known_rows_without_importing_blocked_sales(db):
    wb = openpyxl.load_workbook(workbook())
    ws = wb.active
    extra = [c.value for c in ws[3]]
    extra[0] = "UNKNOWN"
    ws.append(extra)
    file = io.BytesIO()
    wb.save(file)
    file.seek(0)
    p = create_preview(db, SimpleNamespace(username="test"), db.test_shops[0].id, file,
                       2026, date(2026, 8, 30), date(2026, 8, 31), {})
    assert p["status"] == "blocked" and p["error_count"] == 1
    stage = db.get(models.DrsDailyStage, p["stage_id"])
    result = confirm_mappings(db, stage, True)
    assert result["mapped"] == 1 and result["sales_written"] == 0
    assert confirm_mappings(db, stage, True) == result
    assert db.query(models.FactSalesDaily).filter_by(shop_id=db.test_shops[0].id).count() == 0
    assert db.query(models.SalesDailyClaim).filter_by(shop_id=db.test_shops[0].id).count() == 0
    assert stage.status == "blocked" and db.test_products[0].business_sku == "0007"


def test_unknown_historical_cost_is_not_zero_cost_or_profit(db):
    from datetime import timedelta
    from metrics import gross_margin, gross_margin_pct
    from semantic_layer import get_product_performance, get_category_performance
    sku, shop = db.test_products[0].sku_id, db.test_shops[0].id
    db.add(models.FactSalesDaily(sku_id=sku, shop_id=shop, date=date.today() - timedelta(days=1),
                                units_sold=1, revenue=100, cost=60, margin=40))
    unknown = models.FactSalesDaily(sku_id=sku, shop_id=shop, date=date.today(),
                                   units_sold=1, revenue=20, cost=None, margin=None, source="DRS_DAILY")
    db.add(unknown)
    db.commit()
    user = SimpleNamespace(role="admin")
    result = get_product_performance(db, user, sku, days=3, shop_id=shop)
    assert result["revenue"] == 120
    assert result["gross_margin"] is None and result["gross_margin_pct"] is None
    assert result["historical_margin_available"] is False
    category = get_category_performance(db, user, days=3, shop_id=shop)["categories"][0]
    assert category["gross_margin"] is None
    assert gross_margin(100, None) is None and gross_margin_pct(100, None) is None
    # Genuine, explicit zero historical cost remains distinguishable.
    unknown.cost = 0
    db.commit()
    assert get_product_performance(db, user, sku, days=3, shop_id=shop)["gross_margin"] == 60


def test_identity_status_filters_independent_identifiers(db, monkeypatch):
    import authz
    from routes.products import list_products
    monkeypatch.setattr(authz, "user_allowed_shop_ids", lambda *_: {db.test_shops[0].id, db.test_shops[1].id})
    category = "identity-" + uuid4().hex
    a, b = db.test_products
    a.category, b.category = category, category
    a.business_sku, a.ean, a.oem_code = None, "001234", None
    b.business_sku, b.ean, b.oem_code = "0007", "", "OEM-NONUNIQUE"
    db.add_all([
        models.ProductSourceAlias(shop_id=a.shop_id, sku_id=a.sku_id, source="DRS",
                                  code="0007", discriminator=""),
        models.ProductSourceAlias(shop_id=b.shop_id, sku_id=a.sku_id, source="DRS",
                                  code="WRONG-SHOP", discriminator=""),
    ])
    db.commit()
    def query(status):
        return list_products(category=category, subcategory=None, is_private_label=None,
              strategy_type=None, lifecycle_status=None, search=None, identity_status=status,
              hierarchy_node_ids=None, brands=None, statuses=None, sku_query=None,
              page=1, page_size=50, db=db, current_user=SimpleNamespace(id=1))["items"]
    assert len(query("all")) == 2
    assert query("without_business_sku")[0]["source_aliases"] == [
        {"source": "DRS", "code": "0007", "shop_id": a.shop_id}]
    for status in ("with_business_sku", "without_ean", "with_oem"):
        assert [p["sku_id"] for p in query(status)] == [b.sku_id]
    for status in ("without_business_sku", "with_ean"):
        assert [p["sku_id"] for p in query(status)] == [a.sku_id]


def test_identity_labels_distinguish_business_and_technical_fallback(db):
    from services.product_scope import product_label
    p = db.test_products[0]
    assert product_label(p).startswith(f"SKU {p.sku_id} – ")
    assert "zakelijke SKU ontbreekt" in product_label(p)
    p.business_sku = "0007"
    assert product_label(p).startswith("SKU 0007 – ")
    assert f"bron-ID {p.sku_id}" in product_label(p)


def test_feed_overlong_business_source_code_is_rejected_not_truncated(db):
    from feed_ingest import _identity_rows
    p = db.test_products[0]
    with pytest.raises(ValueError, match="niet afgekapt"):
        _identity_rows(db, "priceshape", p.shop_id,
                       [{"sku_id": p.sku_id, "ean": p.ean, "product_group": "x" * 256}], [])
    db.rollback()
    assert db.query(models.ProductSourceAlias).filter_by(shop_id=p.shop_id).count() == 0


@pytest.mark.parametrize("mode", ["aanvullen", "updaten", "vervangen"])
def test_catalog_duplicate_codes_different_eans_block_before_any_writes(db, monkeypatch, mode):
    import asyncio
    from fastapi import UploadFile
    from routes import drs
    from sqlalchemy import event
    shop = db.test_shops[0].id
    first = db.test_products[0]
    second = models.DimProduct(sku_id="duplicate-" + uuid4().hex, shop_id=shop,
                              product_name="Second untouched", product_group="0007", ean="009999",
                              business_sku="0007", is_active=True, import_source="DRS")
    first.business_sku = "0007"
    first.import_source = "DRS"
    db.add(second)
    db.flush()
    db.add_all([models.ProductSourceAlias(shop_id=shop, source="DRS", code="0007",
                                         discriminator=p.ean, sku_id=p.sku_id) for p in [first, second]])
    db.commit()
    before = [(p.sku_id, p.product_name, p.business_sku, p.ean, p.is_active) for p in [first, second]]
    aliases_before = [(a.id, a.sku_id, a.discriminator) for a in
                      db.query(models.ProductSourceAlias).filter_by(shop_id=shop).order_by(models.ProductSourceAlias.id)]
    rows = [{"productcode": "0007", "ean": p.ean, "titel": "Must not overwrite",
             "inkoopprijs ex btw": "1", "verkoopprijs excl btw": "2"} for p in [first, second]]
    monkeypatch.setattr(drs, "_load_drs_xlsx", lambda *_: (
        rows, [{"productcode": "0007", "path": "Must not create"}], []))
    monkeypatch.setattr("services.product_identity.user_allowed_shop_ids", lambda *_: {shop})
    user = SimpleNamespace(username="duplicate-code-test", role="admin")
    writes = []
    def capture(_connection, _cursor, statement, *_args):
        if statement.lstrip().split(" ", 1)[0].upper() in {"INSERT", "UPDATE", "DELETE"}:
            writes.append(statement)
    event.listen(db.bind, "before_cursor_execute", capture)
    try:
        validation = asyncio.run(drs.validate_drs(
            UploadFile(filename="catalog.xlsx", file=io.BytesIO(b"")), db, user))
        assert validation["kan_importeren"] is False
        assert validation["duplicate_skus"] == ["0007"]
        assert any(e["type"] == "dubbele_productcode" and "[2, 3]" in e["bericht"]
                   for e in validation["errors"])
        with pytest.raises(HTTPException) as exc:
            asyncio.run(drs.import_drs(
                UploadFile(filename="catalog.xlsx", file=io.BytesIO(b"")), mode, shop,
                drs.DRS_REPLACE_CONFIRM_PHRASE, db, user))
        assert exc.value.status_code == 422
        assert "0007" in exc.value.detail and "[2, 3]" in exc.value.detail
        assert writes == []
        db.expire_all()
        assert [(p.sku_id, p.product_name, p.business_sku, p.ean, p.is_active)
                for p in [first, second]] == before
        assert [(a.id, a.sku_id, a.discriminator) for a in db.query(
            models.ProductSourceAlias).filter_by(shop_id=shop).order_by(models.ProductSourceAlias.id)] == aliases_before
    finally:
        event.remove(db.bind, "before_cursor_execute", capture)


@pytest.mark.parametrize("zero", [False, True])
def test_zip_sales_respects_daily_claims_and_authorized_product_shop(db, monkeypatch, zero):
    import asyncio
    import zipfile
    from fastapi import UploadFile
    from routes.uploads import upload_bulk
    import services.product_identity as identity
    # Use the second shop to detect accidental default-shop assignment.
    confirm(db, preview(db, shop=1, values=[(0, 0, 0)] * 2 if zero else None))
    shop, sku = db.test_shops[1].id, db.test_products[1].sku_id
    monkeypatch.setattr(identity, "user_allowed_shop_ids", lambda *_: {shop})
    user = SimpleNamespace(username="zip-" + uuid4().hex, role="pricing_manager")
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, "w") as z:
        z.writestr("sales.csv", "sku_id,date,units_sold,revenue\n"
                   f"{sku},2026-08-30,9,90\n{sku},2026-09-01,3,30\n")
    try:
        result = asyncio.run(upload_bulk(
            UploadFile(filename="sales.zip", file=io.BytesIO(raw.getvalue())), db, user))
        assert result["totaal_verwerkt"] == 1 and result["totaal_mislukt"] == 1
        assert any("Overlap" in w for w in result["resultaten"][0]["warnings"])
        fact = db.query(models.FactSalesDaily).filter_by(sku_id=sku, date=date(2026, 9, 1)).one()
        assert fact.shop_id == shop and fact.units_sold == 3
        assert db.query(models.SalesDailyClaim).filter_by(
            shop_id=shop, sku_id=sku, date=date(2026, 9, 1)).one().source == "LEGACY_SALES"
        monkeypatch.setattr(identity, "user_allowed_shop_ids", lambda *_: {db.test_shops[0].id})
        with pytest.raises(HTTPException) as exc:
            asyncio.run(upload_bulk(
                UploadFile(filename="sales.zip", file=io.BytesIO(raw.getvalue())), db, user))
        assert exc.value.status_code == 403
        assert db.query(models.FactSalesDaily).filter_by(sku_id=sku, date=date(2026, 9, 1)).count() == 1
    finally:
        db.rollback()
        db.query(models.ImportJob).filter_by(uploaded_by=user.username).delete()
        db.commit()


@pytest.mark.parametrize("cost,expected", [(None, None), (0, 120), (60, 60)])
def test_bigquery_historical_null_cost_and_explicit_zero_match_postgres(db, monkeypatch, cost, expected):
    import semantic_layer as semantic
    import bigquery.client as client
    import bigquery.cost_tracker as tracker
    queries = []
    class FakeBQ:
        def is_available(self):
            return True
        def query(self, sql, params):
            queries.append(sql)
            if "GROUP BY" in sql:
                assert "COUNT(s.cost) = COUNT(*)" in sql
                assert "p.shop_id = @shop_id" in sql
                rows = [{"category": "Test", "revenue": 120, "units": 2, "cost": cost}]
            else:
                assert "COUNT(cost) = COUNT(*)" in sql
                rows = [{"revenue": 120, "units_sold": 2, "records": 2, "cost": cost}]
            return rows, {"ok": True, "bytes_scanned": 0, "duration_ms": 0}
    monkeypatch.setattr(client, "BigQueryClientWrapper", FakeBQ)
    monkeypatch.setattr(tracker, "record_query_cost", lambda *_: None)
    monkeypatch.setattr(semantic, "_get_bq_project_dataset", lambda: ("project", "dataset"))
    user, shop = SimpleNamespace(role="admin"), db.test_shops[0].id
    product = semantic.get_product_performance(db, user, db.test_products[0].sku_id, days=120, shop_id=shop)
    categories = semantic.get_category_performance(db, user, days=120, shop_id=shop)
    assert product["data_source"] == categories["data_source"] == "bigquery"
    for result in [product, categories["categories"][0]]:
        assert result["gross_margin"] == expected
        assert result["historical_margin_available"] is (cost is not None)
        assert result["gross_margin_pct"] == (expected / 120 * 100 if expected is not None else None)
    assert len(queries) == 2


@pytest.mark.skipif(os.environ.get("RUN_DRS_BENCHMARK") != "1",
                    reason="Explicit isolated-DB full-size performance benchmark")
def test_full_size_confirmation_isolated_benchmark(db):
    """Worst case: all 353,216 product-days nonzero; never uses app/user data."""
    from datetime import datetime, timedelta
    import time
    from psycopg2.extensions import parse_dsn
    assert parse_dsn(os.environ["DATABASE_URL"])["dbname"] != parse_dsn(os.environ["APP_DATABASE_URL_ORIGINAL"])["dbname"]
    shop = db.test_shops[0].id
    prefix = uuid4().hex[:12]
    count = 11038
    products = [{"sku_id": f"bench-{prefix}-{i}", "shop_id": shop,
                 "business_sku": f"B{i:05}", "ean": f"E{i:05}",
                 "product_name": "Synthetic performance product"} for i in range(count)]
    for offset in range(0, count, 1000):
        db.bulk_insert_mappings(models.DimProduct, products[offset:offset + 1000])
    days = [(date(2026, 8, 30) + timedelta(days=i)).isoformat() for i in range(32)]
    stage_id = str(uuid4())
    stage = models.DrsDailyStage(id=stage_id, shop_id=shop, created_by="test",
                                expires_at=datetime.utcnow() + timedelta(hours=1), status="ready",
                                preview={"dates": days, "mapping_proposal_count": count})
    db.add(stage)
    db.flush()
    for offset in range(0, count, 100):
        rows = [{"row": i + 3, "sku_id": p["sku_id"], "code": p["business_sku"],
                 "ean": p["ean"], "values": [[1, "1.21", "1.00"] for _ in days]}
                for i, p in enumerate(products[offset:offset + 100], offset)]
        db.add(models.DrsDailyChunk(stage_id=stage_id, rows=rows))
    db.commit()
    started = time.monotonic()
    result = confirm_stage(db, stage, True)
    duration = time.monotonic() - started
    assert result["inserted"] == count * 32
    assert db.query(models.SalesDailyClaim).filter_by(shop_id=shop).count() == count * 32
    assert db.query(models.FactSalesDaily).filter_by(shop_id=shop).count() == count * 32
    repeat_started = time.monotonic()
    assert confirm_stage(db, stage, True) == result
    print(f"DRS_BENCHMARK products={count} days=32 facts={count * 32} "
          f"confirm_seconds={duration:.3f} retry_seconds={time.monotonic() - repeat_started:.3f}")