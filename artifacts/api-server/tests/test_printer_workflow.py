"""Isolated SQLite workflow; invoke pytest --noconftest to avoid application DB."""
import sys
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from routes import channelmanager as shared
from routes import printer_workflow as api
from services.printer_workflow import (
    PrinterConfig, calculate, target_price, price_index, reconcile, workbook, SHEETS)
from services.channelmanager import number
from models import DimShop, AuditLog
from channel_models import ChannelSnapshot

CONFIG = {"currency": "EUR", "price_basis": "gross", "vat_rate": ".21",
          "active": True, "mapping_confirmed": True, "shops": {"123inkt": 1, "Tonershop": 2}}
PRE = [{"sku": "000001", "shop_id": i, "price": "50"} for i in (1, 2)]
ADVICE = [{"productcode": "000001", "advice": "50"}]
BASE = "/api/channelmanager/printer/1"


@pytest.fixture
def client(monkeypatch):
    engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})
    for table in (DimShop.__table__, AuditLog.__table__, ChannelSnapshot.__table__):
        table.create(engine)
    factory = sessionmaker(bind=engine)
    with factory() as db:
        db.add_all([DimShop(id=i, shop_code=str(i), name=str(i)) for i in (1, 2, 3)])
        db.commit()
    monkeypatch.setattr(shared, "effective_permissions", lambda db, user: {"pricing.view", "pricing.propose", "pricing.approve"})
    monkeypatch.setattr(shared, "user_allowed_shop_ids", lambda db, user: {1, 2})
    state = {"version": 8, "active": True, "rows": ADVICE}
    def feed(db, shop, data):
        if not state["active"] or data["version_id"] != state["version"]:
            raise HTTPException(409, "Inactive or stale advice version")
        return SimpleNamespace(id=state["version"]), state["rows"]
    monkeypatch.setattr(api, "feed_records", feed)
    objects = {}
    def put(content, shop, profile, suffix):
        key = f"private/{len(objects)}.{suffix}"
        objects[key] = content
        return key
    monkeypatch.setattr(api.storage, "put", put)
    monkeypatch.setattr(api.storage, "get", lambda key: objects[key])
    app = FastAPI()
    app.include_router(api.router)
    def session():
        with factory() as db:
            yield db
    app.dependency_overrides[api.get_db] = session
    app.dependency_overrides[api.get_current_user] = lambda: SimpleNamespace(
        id=1, username="manager", is_active=True)
    with TestClient(app) as c:
        c.feed_state, c.objects, c.factory = state, objects, factory
        yield c
    engine.dispose()


def post(client, path, body=None, status=200):
    res = client.post(BASE + path, json={"notes": "explicit review", **(body or {})})
    assert res.status_code == status, res.text
    return res.json()


def setup(client, config=None):
    res = client.put(BASE + "/config", json={
        "expected_revision": 0, "notes": "confirmed scope", "config": config or CONFIG})
    assert res.status_code == 200, res.text
    post(client, "/pre", {"rows": PRE})
    advice = post(client, "/advice", {"profile_id": 3, "version_id": 8, "semantics_confirmed": True})
    return post(client, "/calculate", {"advice_snapshot_id": advice["id"]})


@pytest.mark.parametrize("p,pl,pa,sn,ton", [
    ("19.99", "25.99", "21.49", "22.99", "26.99"),
    ("20", "26.00", "21.50", "23.00", "26.99"),
    ("49.99", "55.99", "51.49", "52.99", "59.99"),
    ("50", "55.00", "50.50", "52.00", "57.99"),
    ("33", "39.00", "34.50", "36.00", "41.99"),
])
def test_excel_boundaries(p, pl, pa, sn, ton):
    assert [target_price(label, number(p)) for label in
            ("Printerland", "Printabout", "Sneltoner", "Tonershop")] == [pl, pa, sn, ton]


def test_missing_zero_duplicates_and_strict_threshold():
    cfg = PrinterConfig(**CONFIG)
    result = calculate(cfg, [{"productcode": "000001", "advice": "60"}], PRE)
    assert not result["warnings"]  # exactly 20%, strictly greater only
    result = calculate(cfg, [{"productcode": "000001", "advice": "60.01"}], PRE)
    assert result["warnings"]
    result = calculate(cfg, ADVICE, PRE[:1])
    assert result["blocks"][0]["reason"] == "missing_target_pre_price"
    result = calculate(cfg, ADVICE, [PRE[0], {**PRE[1], "price": "0"}])
    assert not result["blocks"] and result["rows"][1]["target_price"] is None
    result = calculate(cfg, ADVICE, [{**PRE[0], "price": "0"}, PRE[1]])
    assert result["blocks"][0]["reason"] == "zero_basis_comparison_price"
    result = calculate(cfg, [{"productcode": "000001", "advice": "19.99"}], PRE)
    assert result["blocks"][0]["reason"] == "advice_below_20_do_not_adjust"
    assert price_index(PRE + PRE, {1, 2})[1] == 2
    with pytest.raises(ValueError, match="Conflicting"):
        price_index(PRE + [{**PRE[0], "price": "51"}], {1, 2})
    with pytest.raises(ValueError, match="Conflicting"):
        calculate(cfg, ADVICE + [{"productcode": "000001", "advice": "51"}], PRE)


def test_full_export_and_checked_copy(client):
    calc = setup(client)
    approved = post(client, f"/approve/{calc['id']}")
    exported = post(client, f"/export/{approved['id']}")
    assert post(client, f"/export/{approved['id']}")["id"] == exported["id"]
    assert len(client.objects) == 1
    initial = client.get(BASE + f"/download/{exported['id']}")
    assert initial.status_code == 200
    book = load_workbook(BytesIO(initial.content))
    assert book.sheetnames == list(SHEETS)
    assert "unknown_not_read_back" in [c.value for row in book[SHEETS[4]] for c in row]
    expected = [{"sku": r["sku"], "shop_id": r["shop_id"], "price": r["target_price"]}
                for r in exported["data"]["rows"]]
    partial = post(client, f"/readback/{exported['id']}", {"rows": expected[:1]})
    assert partial["status"] == "mismatch"
    checked = post(client, f"/readback/{exported['id']}", {"rows": expected})
    assert checked["status"] == "confirmed"
    assert checked["data"]["export_sha256"] == exported["data"]["sha256"]
    output = client.get(BASE + f"/readback/{checked['id']}/download")
    assert output.status_code == 200
    book = load_workbook(BytesIO(output.content))
    assert "matched" in [c.value for row in book[SHEETS[4]] for c in row]
    assert client.get(BASE + f"/download/{exported['id']}").content == initial.content
    with client.factory() as db:
        assert db.query(AuditLog).filter(AuditLog.action == "printer_download").count() == 2


def test_stale_and_inactive_sources(client):
    calc = setup(client)
    approved = post(client, f"/approve/{calc['id']}")
    client.feed_state["version"] = 9
    post(client, f"/export/{approved['id']}", status=409)
    client.feed_state["version"] = 8
    post(client, "/pre", {"rows": PRE})
    post(client, f"/approve/{calc['id']}", status=409)
    post(client, f"/export/{approved['id']}", status=409)
    client.feed_state["active"] = False
    post(client, "/advice", {"profile_id": 3, "version_id": 8, "semantics_confirmed": True}, status=409)


def test_permissions_mapping_and_config_guard(client, monkeypatch):
    calc = setup(client)
    approved = post(client, f"/approve/{calc['id']}")
    exported = post(client, f"/export/{approved['id']}")
    monkeypatch.setattr(shared, "user_allowed_shop_ids", lambda db, user: {1})
    assert client.get(BASE).status_code == 404
    assert client.get(BASE + f"/download/{exported['id']}").status_code == 404
    post(client, f"/readback/{exported['id']}", {"rows": PRE}, status=404)


def test_latest_not_lost_by_history_pagination(client):
    setup(client)
    for _ in range(55):
        post(client, "/pre", {"rows": PRE})
    page = client.get(BASE + "?limit=10").json()
    assert page["config"] is not None and page["advice"] is not None
    assert len(page["history"]) == 10 and page["has_more"]
    older = client.get(BASE + f"?limit=10&before_id={page['next_before_id']}").json()
    assert max(r["id"] for r in older["history"]) < min(r["id"] for r in page["history"])


def test_formula_text_not_executable():
    cfg = PrinterConfig(**CONFIG)
    pre = [{**r, "sku": "=unsafe"} for r in PRE]
    advice = [{"productcode": "=unsafe", "advice": "50"}]
    calc = calculate(cfg, advice, pre)
    output = workbook(cfg, calc, {"note": "=danger"}, pre, advice)
    book = load_workbook(BytesIO(output))
    assert not any(c.data_type == "f" for sheet in book for row in sheet for c in row)


def test_readback_suppressed_zero_and_unchanged_products_are_outside_scope():
    exported = [
        {"sku": "001", "shop_id": 1, "target_price": "30"},
        {"sku": "002", "shop_id": 1, "target_price": None},
    ]
    actual = [{"sku": "001", "shop_id": 1, "price": "30"},
              {"sku": "002", "shop_id": 1, "price": "0"},
              {"sku": "003", "shop_id": 1, "price": "99"}]
    result = reconcile(exported, actual, {1})
    assert result["confirmed"] and len(result["extra_rows"]) == 2
    assert len(result["checks"]) == 1
    assert not reconcile(exported, actual[1:], {1})["confirmed"]
    actual[0]["price"] = "31"
    assert not reconcile(exported, actual, {1})["confirmed"]


def test_readback_extra_and_wrong_shop():
    calc = calculate(PrinterConfig(**CONFIG), ADVICE, PRE)
    expected = [{"sku": r["sku"], "shop_id": r["shop_id"], "price": r["target_price"]}
                for r in calc["rows"]]
    assert reconcile(calc["rows"], expected, {1, 2})["confirmed"]
    full = reconcile(calc["rows"], expected + [{"sku": "other", "shop_id": 1, "price": "50"}], {1, 2})
    assert full["confirmed"]
    assert full["extra_rows"] == [{"sku": "other", "shop_id": 1}]
    with pytest.raises(ValueError, match="outside"):
        reconcile(calc["rows"], [{"sku": "000001", "shop_id": 3, "price": "50"}], {1, 2})


def test_config_revision_invalidates_and_default_inactive(client):
    calc = setup(client)
    approved = post(client, f"/approve/{calc['id']}")
    res = client.put(BASE + "/config", json={"expected_revision": 1, "notes": "pause",
                                           "config": {**CONFIG, "active": False}})
    assert res.status_code == 200
    post(client, f"/export/{approved['id']}", status=409)
    post(client, f"/approve/{calc['id']}", status=409)
    assert not PrinterConfig(**{k: v for k, v in CONFIG.items() if k != "active"}).active
    assert client.put(BASE + "/config", json={"expected_revision": 1, "notes": "stale",
                                             "config": CONFIG}).status_code == 409


def test_warning_acknowledgement_and_permissions(client, monkeypatch):
    client.feed_state["rows"] = [{"productcode": "000001", "advice": "61"}]
    calc = setup(client)
    post(client, f"/approve/{calc['id']}", status=409)
    post(client, f"/approve/{calc['id']}", {"acknowledge_warnings": True})
    monkeypatch.setattr(shared, "effective_permissions", lambda db, user: {"pricing.view"})
    post(client, f"/approve/{calc['id']}", {"acknowledge_warnings": True}, status=403)


def test_configurable_rules_and_transform():
    cfg = PrinterConfig(**{**CONFIG, "id_transform": "nl_six_digits", "transformation_confirmed": True})
    pre = [{**r, "sku": "nl000001"} for r in PRE]
    calc = calculate(cfg, [{"productcode": "1", "advice": "50"}], pre)
    assert calc["rows"][0]["source_id"] == "1"
    assert calc["rows"][0]["sku"] == "nl000001"
    with pytest.raises(ValueError, match="collision"):
        calculate(cfg, [{"productcode": "1", "advice": "50"}, {"productcode": "000001", "advice": "50"}], pre)
    cfg.rules["Tonershop"].threshold = "51"
    assert target_price("Tonershop", number("50"), cfg.rules) == "59.99"
