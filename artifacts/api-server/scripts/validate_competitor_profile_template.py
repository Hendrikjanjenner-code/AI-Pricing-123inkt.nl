"""Geautomatiseerde structurele controle voor het concurrentprofiel-template."""

import sys
from pathlib import Path

from openpyxl import load_workbook

from create_competitor_profile_template import CHOICES, OUTPUT, SCHEMA_VERSION, SHEETS


def validate(path: Path) -> None:
    wb = load_workbook(path, data_only=False)
    required = ["Instructies", "Bedrijven", "Relaties_Lidmaatschappen", "Bronnen"]
    assert wb.sheetnames == required, f"Bladen wijken af: {wb.sheetnames}"

    for sheet, headers in SHEETS.items():
        ws = wb[sheet]
        actual = [cell.value for cell in ws[1]]
        assert actual == headers, f"Headers wijken af in {sheet}"
        assert ws.freeze_panes == "A2", f"Kopregel niet vastgezet in {sheet}"
        assert ws.tables, f"Filtertabel ontbreekt in {sheet}"
        versions = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
        assert versions == {SCHEMA_VERSION}, f"Verkeerde schema_version in {sheet}"

        validations = list(ws.data_validations.dataValidation)
        for field in set(headers).intersection(CHOICES):
            expected = '"' + ",".join(CHOICES[field]) + '"'
            assert any(v.type == "list" and v.formula1 == expected for v in validations), (
                f"Keuzelijst ontbreekt voor {sheet}.{field}"
            )

    for sheet, key in [
        ("Bedrijven", ("company_key", "field_key")),
        ("Relaties_Lidmaatschappen", ("relation_key",)),
        ("Bronnen", ("source_id",)),
    ]:
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        indexes = [headers.index(name) for name in key]
        values = [tuple(row[i] for i in indexes) for row in ws.iter_rows(min_row=2, values_only=True)]
        assert all(all(v not in (None, "") for v in item) for item in values), f"Lege sleutel in {sheet}"
        assert len(values) == len(set(values)), f"Dubbele sleutel in {sheet}"

    company_ws = wb["Bedrijven"]
    headers = [c.value for c in company_ws[1]]
    for row in company_ws.iter_rows(min_row=2, values_only=True):
        item = dict(zip(headers, row))
        value_count = sum(item[name] is not None for name in
                          ("value_text", "value_number", "value_boolean", "value_date"))
        if item["status"] == "unknown":
            assert value_count == 0, "Unknown bevat toch een waarde"
        else:
            assert value_count == 1, "Known moet exact één value_* veld bevatten"
        assert 0 <= item["confidence"] <= 1

    source_ids = {r[1] for r in wb["Bronnen"].iter_rows(min_row=2, values_only=True)}
    for sheet in ("Bedrijven", "Relaties_Lidmaatschappen"):
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        source_col = headers.index("source_id")
        for row in ws.iter_rows(min_row=2, values_only=True):
            assert row[source_col] in source_ids, f"Onbekende source_id in {sheet}"

    relation_ws = wb["Relaties_Lidmaatschappen"]
    headers = [c.value for c in relation_ws[1]]
    for row in relation_ws.iter_rows(min_row=2, values_only=True):
        item = dict(zip(headers, row))
        legal = {"parent", "ultimate_parent", "subsidiary", "sister"}
        assert (item["relation_type"] in legal) == (item["relation_class"] == "juridisch"), (
            "Juridische en commerciële relatieclassificatie niet consistent"
        )

    wb.close()
    print(f"OK: {path} voldoet aan {SCHEMA_VERSION}")


if __name__ == "__main__":
    validate(Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT)